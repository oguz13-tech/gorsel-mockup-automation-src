# -*- coding: utf-8 -*-
"""Independent end-user utility suite for the mockup workspace.

Each tool can be started on its own and never depends on the factory render
loop. Destructive operations are intentionally blocked inside the live
template pools so utility usage cannot invalidate approved template hashes.
"""

from __future__ import annotations

import argparse
import base64
import importlib.util
import json
import os
import re
import sqlite3
import subprocess
import sys
import uuid
import time
import contextlib
from pathlib import Path
from typing import Callable

import requests
from openpyxl import Workbook
from PIL import Image

try:
    from .core.template_identity import (
        TemplateNumberConflictError,
        claim_number_block_once,
        category_for_prefix,
        same_stem_files,
        validate_number_block,
    )
    from .core.path_registry import (
        resolve_runtime_root as resolve_registered_runtime_root,
        runtime_path,
    )
except ImportError:
    from core.template_identity import (
        TemplateNumberConflictError,
        category_for_prefix,
        same_stem_files,
        validate_number_block,
    )
    from core.path_registry import (
        resolve_runtime_root as resolve_registered_runtime_root,
        runtime_path,
    )


SOURCE_ROOT = Path(__file__).resolve().parent.parent
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


def resolve_runtime_root() -> Path:
    return resolve_registered_runtime_root(SOURCE_ROOT)



RUNTIME_ROOT = resolve_runtime_root()
os.environ["MOCKUP_RUNTIME_WORKSPACE"] = str(RUNTIME_ROOT)


def _ui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, simpledialog

    root = tk.Tk()
    root.withdraw()
    try:
        root.attributes("-topmost", True)
    except Exception:
        pass
    return root, filedialog, messagebox, simpledialog


def _is_under(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except Exception:
        return False


def _protected_template_roots(runtime_root: Path | None = None) -> list[Path]:
    root = (runtime_root or RUNTIME_ROOT).resolve()
    return [
        runtime_path("WORKING_POOL", root),
        runtime_path("ORIGINAL_ARCHIVE", root),
    ]


def ensure_not_live_template_path(
    paths: list[Path],
    runtime_root: Path | None = None,
) -> None:
    protected = _protected_template_roots(runtime_root)
    for path in paths:
        for root in protected:
            if root.exists() and _is_under(path, root):
                raise ValueError(
                    "Bu araç aktif şablon veya orijinal-yedek klasöründe "
                    "dosya adı değiştiremez. Bu klasörleri İnceleme Merkezi "
                    "yönetir."
                )


def _sanitize_prefix(prefix: str) -> str:
    prefix = str(prefix or "").strip()
    prefix = re.sub(r'[<>:"/\\|?*]+', "_", prefix)
    prefix = prefix.rstrip(" .")
    if not prefix:
        raise ValueError("Geçerli bir başlangıç adı yazın.")
    return prefix


@contextlib.contextmanager
def _template_rename_lock(runtime_root: Path):
    """Same physical lock used by staging watcher, across processes.

    Fail closed on stale locks; never remove someone else's active lock.
    """
    locks = Path(runtime_root).resolve() / "SABLON_V2_MIMARI" / "recovery" / "locks"
    locks.mkdir(parents=True, exist_ok=True)
    path = locks / "staging_intake.lock"
    deadline = time.monotonic() + 10
    fd = None
    while time.monotonic() < deadline:
        try:
            fd = os.open(str(path), os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            os.write(fd, uuid.uuid4().hex.encode("ascii"))
            break
        except FileExistsError:
            time.sleep(0.05)
    if fd is None:
        raise RuntimeError("NUMBER_CLAIM_LOCK_BUSY_OR_STALE")
    try:
        yield
    finally:
        os.close(fd)
        path.unlink(missing_ok=True)


def rename_images(
    files: list[Path],
    prefix: str,
    runtime_root: Path | None = None,
    start_number: int = 1,
) -> list[Path]:
    """Collision-safe, two-phase rename for selected files in one folder."""
    selected = [Path(p).resolve() for p in files]
    if not selected:
        raise ValueError("İsimlendirilecek görsel seçilmedi.")
    if any(not p.is_file() for p in selected):
        raise ValueError("Seçim içinde bulunamayan dosya var.")
    if any(p.suffix.lower() not in IMAGE_EXTS for p in selected):
        raise ValueError("Seçimde desteklenmeyen dosya türü var.")

    parents = {p.parent for p in selected}
    if len(parents) != 1:
        raise ValueError("Görseller aynı klasörden seçilmelidir.")

    ensure_not_live_template_path(selected, runtime_root)
    prefix = _sanitize_prefix(prefix)
    try:
        start_number = int(start_number)
    except (TypeError, ValueError):
        raise ValueError("Başlangıç numarası geçerli bir sayı olmalıdır.")
    if start_number < 0:
        raise ValueError("Başlangıç numarası 0 veya daha büyük olmalıdır.")

    selected = sorted(selected, key=lambda p: p.name.casefold())
    parent = selected[0].parent
    selected_set = {p.resolve() for p in selected}

    template_category = category_for_prefix(prefix)
    if template_category and start_number < 1:
        raise ValueError("Şablon numarası 1 veya daha büyük olmalıdır.")

    last_number = start_number + len(selected) - 1
    digits = max(3, len(str(last_number)))
    targets = [
        parent
        / (
            f"{prefix}{number:0{digits}d}"
            f"{source.suffix.lower()}"
        )
        for source, number in zip(
            selected,
            range(start_number, start_number + len(selected)),
        )
    ]
    if len({p.name.casefold() for p in targets}) != len(targets):
        raise ValueError("Yeni dosya adları kendi içinde çakışıyor.")

    for target in targets:
        stem_matches = same_stem_files(
            target.parent,
            target.name,
            ignore_paths=selected_set,
        )
        if stem_matches:
            existing = ", ".join(path.name for path in stem_matches)
            existing_exts = sorted({
                path.suffix.casefold()
                for path in stem_matches
            })
            incoming_ext = target.suffix.casefold()
            different_extension = any(
                ext != incoming_ext
                for ext in existing_exts
            )
            headline = (
                "AYNI İSİM, FARKLI UZANTI"
                if different_extension
                else "AYNI İSİM / AYNI UZANTI"
            )
            raise FileExistsError(
                f"{headline}: hedef '{target.stem}' zaten kullanılıyor. "
                f"Yeni hedef={target.name}; mevcut={existing}. "
                "Dosyalar değiştirilmedi."
            )

    # Keep the shared lock across BOTH claims and filesystem publication.
    # The registry is still durable if any disk operation later fails.
    root = Path(runtime_root or RUNTIME_ROOT).resolve()
    with (_template_rename_lock(root) if template_category else contextlib.nullcontext()):
        if template_category:
            claim_number_block_once(
                root, template_category, start_number, len(selected),
                canonical_names=[target.name for target in targets],
                ignore_paths=selected,
            )

        temp_pairs: list[tuple[Path, Path]] = []
        final_pairs: list[tuple[Path, Path]] = []
        try:
            for source in selected:
                temp = source.with_name(
                    f".__jarvis_rename_{uuid.uuid4().hex}{source.suffix.lower()}"
                )
                source.rename(temp)
                temp_pairs.append((source, temp))

            for (source, temp), target in zip(temp_pairs, targets):
                # Exclusive publication: Path.rename can overwrite targets on
                # POSIX after the precheck; hard-link creation cannot.
                os.link(str(temp), str(target))
                final_pairs.append((temp, target))
            # Only remove temporary names once EVERY final target exists.
            for original, temp in temp_pairs:
                temp.unlink()
            return targets
        except Exception:
            # Remove ONLY published files still hard-linked to our temp;
            # never delete another process's later edits.
            for temp, target in reversed(final_pairs):
                try:
                    if temp.exists() and target.exists() and os.path.samefile(temp, target):
                        target.unlink()
                except OSError:
                    pass
            for original, temp in reversed(temp_pairs):
                if temp.exists() and not original.exists():
                    try:
                        temp.rename(original)
                    except OSError:
                        # Keep recoverable .__jarvis_rename temp; never unlink.
                        pass
            raise


def _unique_destination(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 10000):
        candidate = path.with_name(
            f"{path.stem}_{index}{path.suffix}"
        )
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Çıktı adı üretilemedi: {path.name}")


def convert_images(
    files: list[Path],
    target_format: str,
) -> tuple[Path, list[Path]]:
    target = str(target_format or "").upper().strip()
    if target not in {"JPG", "PNG"}:
        raise ValueError("Hedef format JPG veya PNG olmalıdır.")

    selected = [Path(p).resolve() for p in files]
    if not selected:
        raise ValueError("Dönüştürülecek görsel seçilmedi.")
    parents = {p.parent for p in selected}
    if len(parents) != 1:
        raise ValueError("Görseller aynı klasörden seçilmelidir.")

    output_dir = selected[0].parent / f"FORMAT_{target}_CIKTILARI"
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = []

    for source in selected:
        if source.suffix.lower() not in IMAGE_EXTS:
            continue
        with Image.open(source) as image:
            if target == "JPG":
                destination = _unique_destination(
                    output_dir / f"{source.stem}.jpg"
                )
                image.convert("RGB").save(
                    destination,
                    format="JPEG",
                    quality=95,
                    optimize=True,
                )
            else:
                destination = _unique_destination(
                    output_dir / f"{source.stem}.png"
                )
                image.convert("RGBA").save(
                    destination,
                    format="PNG",
                    optimize=True,
                )
            outputs.append(destination)

    return output_dir, outputs


def convert_folder_to_png_recursive(
    source_dir: Path,
    *,
    output_dir: Path | None = None,
    runtime_root: Path | None = None,
    progress: Callable[[int, int, str], None] | None = None,
) -> dict:
    """Recursively convert non-PNG images while preserving subfolders.

    Safety contract:
    - source files are never deleted or overwritten;
    - active template/original-archive trees are blocked;
    - existing PNG source files are skipped;
    - existing destination files are never overwritten;
    - every failure is returned in the final report.
    """
    source = Path(source_dir).expanduser().resolve()
    if not source.is_dir():
        raise ValueError("Kaynak klasör bulunamadı.")

    ensure_not_live_template_path([source], runtime_root)

    if output_dir is None:
        destination_root = source.parent / f"{source.name}_PNG_CIKTILARI"
    else:
        destination_root = Path(output_dir).expanduser().resolve()

    try:
        destination_root.relative_to(source)
        raise ValueError(
            "Çıktı klasörü kaynak klasörün içinde olamaz; recursive döngü engellendi."
        )
    except ValueError as exc:
        if "recursive döngü" in str(exc):
            raise

    destination_root.mkdir(parents=True, exist_ok=True)

    report = {
        "source_dir": str(source),
        "output_dir": str(destination_root),
        "scanned_image_count": 0,
        "converted_count": 0,
        "skipped_png_count": 0,
        "skipped_existing_count": 0,
        "failed_count": 0,
        "converted": [],
        "skipped_png": [],
        "skipped_existing": [],
        "failed": [],
    }

    candidates = sorted(
        (
            path for path in source.rglob("*")
            if path.is_file() and path.suffix.lower() in IMAGE_EXTS
        ),
        key=lambda path: path.as_posix().casefold(),
    )

    total = len(candidates)
    if progress:
        progress(0, total, "Görseller tarandı; dönüşüm başlıyor.")

    for index, source_file in enumerate(candidates, start=1):
        report["scanned_image_count"] += 1
        relative = source_file.relative_to(source)
        if progress:
            progress(
                index - 1,
                total,
                f"İşleniyor: {relative.as_posix()}",
            )

        if source_file.suffix.lower() == ".png":
            report["skipped_png_count"] += 1
            report["skipped_png"].append(relative.as_posix())
            continue

        destination = (
            destination_root
            / relative.parent
            / f"{source_file.stem}.png"
        )
        if destination.exists():
            report["skipped_existing_count"] += 1
            report["skipped_existing"].append(
                destination.relative_to(destination_root).as_posix()
            )
            continue

        try:
            destination.parent.mkdir(parents=True, exist_ok=True)
            with Image.open(source_file) as image:
                image.convert("RGBA").save(
                    destination,
                    format="PNG",
                    optimize=True,
                )
            report["converted_count"] += 1
            report["converted"].append({
                "source": relative.as_posix(),
                "output": destination.relative_to(
                    destination_root
                ).as_posix(),
            })
        except Exception as exc:
            report["failed_count"] += 1
            report["failed"].append({
                "source": relative.as_posix(),
                "error": str(exc),
            })

    if progress:
        progress(
            total,
            total,
            "Recursive PNG dönüşümü tamamlandı.",
        )
    return report


def watermark_images(
    logo_path: Path,
    files: list[Path],
) -> tuple[Path, list[Path]]:
    logo_path = Path(logo_path).resolve()
    selected = [Path(p).resolve() for p in files]
    if not logo_path.is_file():
        raise ValueError("Filigran logosu bulunamadı.")
    if not selected:
        raise ValueError("Filigran uygulanacak görsel seçilmedi.")

    parents = {p.parent for p in selected}
    if len(parents) != 1:
        raise ValueError("Görseller aynı klasörden seçilmelidir.")

    output_dir = selected[0].parent / "FILIGRANLI_GORSELLER"
    output_dir.mkdir(parents=True, exist_ok=True)
    outputs = []

    with Image.open(logo_path) as logo_source:
        logo = logo_source.convert("RGBA")
        for source in selected:
            with Image.open(source) as input_image:
                image = input_image.convert("RGBA")
                wm_width = max(1, int(image.width * 0.20))
                scale = wm_width / max(1, logo.width)
                wm_height = max(1, int(logo.height * scale))
                resized = logo.resize(
                    (wm_width, wm_height),
                    Image.Resampling.LANCZOS,
                )
                pad_x = int(image.width * 0.05)
                pad_y = int(image.height * 0.05)
                position = (
                    max(0, image.width - wm_width - pad_x),
                    max(0, image.height - wm_height - pad_y),
                )
                composed = image.copy()
                composed.alpha_composite(resized, position)

                destination = _unique_destination(
                    output_dir / source.name
                )
                if destination.suffix.lower() in {".jpg", ".jpeg"}:
                    composed.convert("RGB").save(
                        destination,
                        quality=95,
                        optimize=True,
                    )
                else:
                    composed.save(destination)
                outputs.append(destination)

    return output_dir, outputs


def _load_cloud_settings() -> dict:
    db_path = RUNTIME_ROOT / "SABLON_V2_MIMARI" / "ayarlar.db"
    if db_path.is_file():
        try:
            conn = sqlite3.connect(str(db_path))
            row = conn.execute(
                "SELECT value FROM settings WHERE key='bulut_ayarlari'"
            ).fetchone()
            conn.close()
            if row and row[0]:
                value = json.loads(row[0])
                if isinstance(value, dict):
                    return value
        except Exception:
            pass

    for candidate in (
        RUNTIME_ROOT / "ayarlar.json",
        SOURCE_ROOT / "ayarlar.json",
    ):
        if not candidate.is_file():
            continue
        try:
            data = json.loads(candidate.read_text(encoding="utf-8"))
            cfg = data.get("bulut_ayarlari", {})
            if isinstance(cfg, dict):
                return cfg
        except Exception:
            pass
    return {}


def _imgbb_keys() -> list[str]:
    cfg = _load_cloud_settings()
    configured = cfg.get("imgbb_keys") or []
    if not isinstance(configured, list):
        configured = []
    env = [
        key.strip()
        for key in os.environ.get("IMGBB_KEYS", "").split(",")
        if key.strip()
    ]
    return [
        str(key).strip()
        for key in (configured or env)
        if str(key).strip()
    ]


def upload_to_imgbb(image_path: Path, api_keys: list[str]) -> str | None:
    encoded = base64.b64encode(image_path.read_bytes()).decode("ascii")
    for key in api_keys:
        try:
            response = requests.post(
                "https://api.imgbb.com/1/upload",
                data={"key": key, "image": encoded},
                timeout=45,
            )
            if response.status_code == 200:
                payload = response.json()
                url = (
                    payload.get("data", {}).get("url")
                    if isinstance(payload, dict)
                    else None
                )
                if url:
                    return str(url)
        except Exception:
            continue
    return None


def images_to_link_excel(files: list[Path]) -> tuple[Path, int, int]:
    selected = [Path(p).resolve() for p in files]
    if not selected:
        raise ValueError("Linke dönüştürülecek görsel seçilmedi.")
    parents = {p.parent for p in selected}
    if len(parents) != 1:
        raise ValueError("Görseller aynı klasörden seçilmelidir.")

    keys = _imgbb_keys()
    if not keys:
        raise RuntimeError(
            "ImgBB anahtarı bulunamadı. Sistem ayarlarında "
            "bulut_ayarlari.imgbb_keys yapılandırılmalı."
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Görsel Linkleri"
    sheet.append(["Dosya Adı", "Resim Linki", "Sistem Dosya Yolu"])

    success = 0
    failed = 0
    for index, source in enumerate(selected, start=1):
        print(f"[{index}/{len(selected)}] Yükleniyor: {source.name}")
        url = upload_to_imgbb(source, keys)
        if url:
            success += 1
            sheet.append([source.name, url, str(source)])
        else:
            failed += 1
            sheet.append([source.name, "YÜKLEME_HATASI", str(source)])

    output = _unique_destination(
        selected[0].parent / "Gorsel_Linkleri.xlsx"
    )
    workbook.save(output)
    return output, success, failed


def _load_seo_module():
    script = (
        SOURCE_ROOT
        / "ÜRÜN ADI YAZDIRMA BOTU"
        / "Benzersiz-ürün-adlari-yaz.py"
    )
    if not script.is_file():
        raise FileNotFoundError("SEO üretici scripti bulunamadı.")
    spec = importlib.util.spec_from_file_location("jarvis_seo_tool", script)
    if spec is None or spec.loader is None:
        raise RuntimeError("SEO üretici yüklenemedi.")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def generate_seo_excel(
    count: int,
    mode: str,
    output_dir: Path,
) -> Path:
    count = max(1, min(int(count), 5000))
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    module = _load_seo_module()

    if mode == "titles":
        groups = module.urun_isim_dosyasini_oku()
        values = module.isimleri_uret(count, groups)
        if not values:
            raise RuntimeError("SEO ürün adı üretilemedi.")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SEO Ürün Adları"
        sheet.append(["Sıra", "Ürün Adı", "Karakter Sayısı"])
        for index, value in enumerate(values, start=1):
            sheet.append([index, value, len(value)])
        output = _unique_destination(
            output_dir / "SEO_Urun_Adlari.xlsx"
        )
    elif mode == "descriptions":
        sentences, keywords = module.aciklama_dosyasini_oku()
        values = module.aciklamalari_uret(
            count,
            sentences,
            keywords,
        )
        if not values:
            raise RuntimeError("SEO ürün açıklaması üretilemedi.")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "SEO Açıklamalar"
        sheet.append(["Sıra", "Ürün Açıklaması"])
        for index, value in enumerate(values, start=1):
            sheet.append([index, value])
        output = _unique_destination(
            output_dir / "SEO_Urun_Aciklamalari.xlsx"
        )
    else:
        raise ValueError("Bilinmeyen SEO modu.")

    workbook.save(output)
    return output


def _open_folder(path: Path) -> None:
    try:
        if os.name == "nt":
            os.startfile(str(path))
    except Exception:
        pass


def interactive_rename() -> int:
    root, filedialog, messagebox, simpledialog = _ui()
    try:
        files = filedialog.askopenfilenames(
            title="İsimlendirilecek Görselleri Seçin",
            filetypes=[
                ("Görseller", "*.png *.jpg *.jpeg *.webp *.bmp"),
                ("Tüm Dosyalar", "*.*"),
            ],
        )
        if not files:
            return 0

        prefix = simpledialog.askstring(
            "Görsel İsimlendirici - 1/2",
            "TEMEL ADI yazın.\n"
            "Sayı yazmayın.\n\n"
            "Örnek: Wifi_Kutusu_",
            initialvalue="Gorsel_",
            parent=root,
        )
        if prefix is None:
            return 0

        suggested_start = 1
        while True:
            start_number = simpledialog.askinteger(
                "Görsel İsimlendirici - 2/2",
                "Kaçtan başlasın?\n\n"
                "Örnek: 135 yazarsanız ilk dosya "
                "Wifi_Kutusu_135 olur.\n\n"
                "Daha önce kullanılmış numaralar tekrar kullanılamaz.",
                initialvalue=suggested_start,
                minvalue=0,
                maxvalue=9999999,
                parent=root,
            )
            if start_number is None:
                return 0

            try:
                outputs = rename_images(
                    [Path(p) for p in files],
                    prefix,
                    RUNTIME_ROOT,
                    start_number=start_number,
                )
                break
            except TemplateNumberConflictError as exc:
                suggested_start = (
                    exc.suggested_start
                    if exc.suggested_start is not None
                    else start_number + 1
                )
                messagebox.showwarning(
                    "Bu Numara Daha Önce Kullanılmış",
                    f"{exc}\n\n"
                    "İşlem yapılmadı. Yeni bir başlangıç numarası seçmeniz "
                    "zorunlu. Araç şimdi tekrar numara soracak.",
                    parent=root,
                )

        _open_folder(outputs[0].parent)
        messagebox.showinfo(
            "Tamamlandı",
            f"{len(outputs)} görsel yeniden adlandırıldı.\n"
            f"İlk dosya: {outputs[0].name}\n"
            f"Son dosya: {outputs[-1].name}",
            parent=root,
        )
        return 0
    except Exception as exc:
        messagebox.showerror("İsimlendirme Hatası", str(exc), parent=root)
        return 1
    finally:
        root.destroy()


def interactive_convert() -> int:
    root, filedialog, messagebox, simpledialog = _ui()
    try:
        files = filedialog.askopenfilenames(
            title="Dönüştürülecek Görselleri Seçin",
            filetypes=[
                ("Görseller", "*.png *.jpg *.jpeg *.webp *.bmp"),
                ("Tüm Dosyalar", "*.*"),
            ],
        )
        if not files:
            return 0
        target = simpledialog.askstring(
            "Format Çevirici",
            "Hedef formatı yazın: JPG veya PNG",
            initialvalue="JPG",
            parent=root,
        )
        if target is None:
            return 0
        output_dir, outputs = convert_images(
            [Path(p) for p in files],
            target,
        )
        _open_folder(output_dir)
        messagebox.showinfo(
            "Tamamlandı",
            f"{len(outputs)} görsel dönüştürüldü.\n"
            f"Çıktı: {output_dir}",
            parent=root,
        )
        return 0
    except Exception as exc:
        messagebox.showerror("Format Hatası", str(exc), parent=root)
        return 1
    finally:
        root.destroy()


def interactive_recursive_png() -> int:
    root, filedialog, messagebox, _simpledialog = _ui()
    progress_window = None
    try:
        source_dir = filedialog.askdirectory(
            title="PNG'ye Recursive Dönüştürülecek Ana Klasörü Seçin"
        )
        if not source_dir:
            return 0

        import tkinter as tk
        from tkinter import ttk

        progress_window = tk.Toplevel(root)
        progress_window.title("Recursive PNG Dönüştürücü v2")
        progress_window.geometry("520x150")
        progress_window.resizable(False, False)
        try:
            progress_window.attributes("-topmost", True)
        except Exception:
            pass

        title = tk.Label(
            progress_window,
            text="Alt klasörler dahil PNG dönüşümü hazırlanıyor...",
            anchor="w",
        )
        title.pack(fill="x", padx=16, pady=(16, 8))

        bar = ttk.Progressbar(
            progress_window,
            orient="horizontal",
            mode="determinate",
            maximum=100,
        )
        bar.pack(fill="x", padx=16, pady=6)

        status = tk.Label(
            progress_window,
            text="",
            anchor="w",
            justify="left",
            wraplength=480,
        )
        status.pack(fill="x", padx=16, pady=(6, 12))

        def on_progress(current: int, total: int, message: str) -> None:
            denominator = max(1, int(total))
            bar["value"] = min(
                100,
                max(0, (int(current) / denominator) * 100),
            )
            title.config(
                text=f"{int(current)} / {int(total)} görsel işlendi"
            )
            status.config(text=str(message or ""))
            progress_window.update_idletasks()
            progress_window.update()

        report = convert_folder_to_png_recursive(
            Path(source_dir),
            runtime_root=RUNTIME_ROOT,
            progress=on_progress,
        )
        output = Path(report["output_dir"])

        report_json = _unique_destination(
            output / "PNG_DONUSUM_RAPORU.json"
        )
        report_json.write_text(
            json.dumps(report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        report_txt = _unique_destination(
            output / "PNG_DONUSUM_RAPORU.txt"
        )
        report_txt.write_text(
            "\n".join([
                "RECURSIVE PNG DONUSUM RAPORU",
                f"Kaynak: {report['source_dir']}",
                f"Cikti: {report['output_dir']}",
                f"Taranan: {report['scanned_image_count']}",
                f"Donusturulen: {report['converted_count']}",
                f"PNG atlanan: {report['skipped_png_count']}",
                f"Mevcut cikti atlanan: {report['skipped_existing_count']}",
                f"Hata: {report['failed_count']}",
                "",
                "HATALAR:",
                *[
                    f"- {item['source']}: {item['error']}"
                    for item in report["failed"]
                ],
            ]),
            encoding="utf-8",
        )

        if progress_window is not None:
            progress_window.destroy()
            progress_window = None

        _open_folder(output)

        failures = report["failed"]
        failure_preview = ""
        if failures:
            failure_preview = "\n\nİlk hatalar:\n" + "\n".join(
                f"- {item['source']}: {item['error']}"
                for item in failures[:10]
            )

        message = (
            "Recursive PNG dönüşümü tamamlandı.\n\n"
            f"Taranan görsel: {report['scanned_image_count']}\n"
            f"Dönüştürülen: {report['converted_count']}\n"
            f"Zaten PNG olduğu için atlanan: {report['skipped_png_count']}\n"
            f"Çıktı zaten var diye atlanan: {report['skipped_existing_count']}\n"
            f"Hata: {report['failed_count']}\n\n"
            f"Çıktı: {output}\n"
            f"JSON rapor: {report_json.name}\n"
            f"Metin raporu: {report_txt.name}"
            + failure_preview
        )

        if failures:
            messagebox.showwarning(
                "Recursive PNG — Kısmi Tamamlandı",
                message,
                parent=root,
            )
            return 1

        messagebox.showinfo(
            "Recursive PNG — Tamamlandı",
            message,
            parent=root,
        )
        return 0
    except Exception as exc:
        if progress_window is not None:
            try:
                progress_window.destroy()
            except Exception:
                pass
        messagebox.showerror(
            "Recursive PNG Hatası",
            str(exc),
            parent=root,
        )
        return 1
    finally:
        root.destroy()


def interactive_watermark() -> int:
    root, filedialog, messagebox, _simpledialog = _ui()
    try:
        logo = filedialog.askopenfilename(
            title="1. Adım: Filigran Logosunu Seçin",
            filetypes=[("PNG Görseli", "*.png"), ("Görseller", "*.*")],
        )
        if not logo:
            return 0
        files = filedialog.askopenfilenames(
            title="2. Adım: Filigran Uygulanacak Görselleri Seçin",
            filetypes=[
                ("Görseller", "*.png *.jpg *.jpeg *.webp"),
                ("Tüm Dosyalar", "*.*"),
            ],
        )
        if not files:
            return 0
        output_dir, outputs = watermark_images(
            Path(logo),
            [Path(p) for p in files],
        )
        _open_folder(output_dir)
        messagebox.showinfo(
            "Tamamlandı",
            f"{len(outputs)} görsele filigran uygulandı.\n"
            f"Çıktı: {output_dir}",
            parent=root,
        )
        return 0
    except Exception as exc:
        messagebox.showerror("Filigran Hatası", str(exc), parent=root)
        return 1
    finally:
        root.destroy()


def interactive_links() -> int:
    root, filedialog, messagebox, _simpledialog = _ui()
    try:
        files = filedialog.askopenfilenames(
            title="Linke Dönüştürülecek Görselleri Seçin",
            filetypes=[
                ("Görseller", "*.png *.jpg *.jpeg *.webp *.bmp"),
                ("Tüm Dosyalar", "*.*"),
            ],
        )
        if not files:
            return 0
        output, success, failed = images_to_link_excel(
            [Path(p) for p in files]
        )
        _open_folder(output.parent)
        if failed:
            messagebox.showwarning(
                "İşlem Tamamlandı",
                f"Başarılı: {success}\nHatalı: {failed}\n"
                f"Excel: {output}",
                parent=root,
            )
            return 1
        messagebox.showinfo(
            "Tamamlandı",
            f"{success} görsel linke dönüştürüldü.\nExcel: {output}",
            parent=root,
        )
        return 0
    except Exception as exc:
        messagebox.showerror("Link Dönüştürme Hatası", str(exc), parent=root)
        return 1
    finally:
        root.destroy()


def interactive_seo(mode: str) -> int:
    root, filedialog, messagebox, simpledialog = _ui()
    try:
        count = simpledialog.askinteger(
            "SEO Üretici",
            "Kaç adet üretilecek?",
            initialvalue=10,
            minvalue=1,
            maxvalue=5000,
            parent=root,
        )
        if not count:
            return 0
        output_dir = filedialog.askdirectory(
            title="Excel Çıktısının Kaydedileceği Klasörü Seçin"
        )
        if not output_dir:
            return 0
        output = generate_seo_excel(
            count,
            mode,
            Path(output_dir),
        )
        _open_folder(output.parent)
        label = (
            "SEO ürün adı"
            if mode == "titles"
            else "SEO ürün açıklaması"
        )
        messagebox.showinfo(
            "Tamamlandı",
            f"{count} hedefiyle {label} üretildi.\nExcel: {output}",
            parent=root,
        )
        return 0
    except Exception as exc:
        messagebox.showerror("SEO Üretim Hatası", str(exc), parent=root)
        return 1
    finally:
        root.destroy()


def interactive_4k() -> int:
    root, filedialog, messagebox, _simpledialog = _ui()
    try:
        input_dir = filedialog.askdirectory(
            title="4K Yapılacak Görsellerin Klasörünü Seçin"
        )
        if not input_dir:
            return 0
        source = Path(input_dir).resolve()
        output = source / "4K_POWER_TOOL_OUTPUT"
        script = (
            SOURCE_ROOT
            / "99_SISTEM_MOTORLARI_VE_KODLAR"
            / "SISTEM_MOTORU"
            / "kalite_artirici_otomasyon.py"
        )
        if not script.is_file():
            raise FileNotFoundError("4K kalite motoru bulunamadı.")

        cmd = [
            sys.executable,
            str(script),
            "--input",
            str(source),
            "--output",
            str(output),
            "--force-ai",
        ]
        kwargs = {"cwd": str(SOURCE_ROOT)}
        if os.name == "nt":
            kwargs["creationflags"] = subprocess.CREATE_NEW_CONSOLE
        subprocess.Popen(cmd, **kwargs)
        messagebox.showinfo(
            "4K İşlemi Başlatıldı",
            "4K kalite işlemi ayrı pencerede çalışıyor.\n"
            f"Çıktı klasörü: {output}",
            parent=root,
        )
        return 0
    except Exception as exc:
        messagebox.showerror("4K Kalite Hatası", str(exc), parent=root)
        return 1
    finally:
        root.destroy()


TOOLS = {
    "rename": interactive_rename,
    "convert": interactive_convert,
    "png_recursive": interactive_recursive_png,
    "watermark": interactive_watermark,
    "4k": interactive_4k,
    "links": interactive_links,
    "seo_titles": lambda: interactive_seo("titles"),
    "seo_descriptions": lambda: interactive_seo("descriptions"),
}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Bağımsız e-ticaret görsel araçları."
    )
    parser.add_argument("tool", choices=sorted(TOOLS))
    args = parser.parse_args()
    return int(TOOLS[args.tool]())


if __name__ == "__main__":
    raise SystemExit(main())
