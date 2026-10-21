# -*- coding: utf-8 -*-
"""PETEK V3 local-only serial-production queue and commerce export.

The old six-category factory remains untouched. This job path uses the
approved-mask store, deterministic artifact identities and an isolated SQLite
ledger. It never uploads, posts listings, modifies originals or overwrites
existing image output. Approval gate is per source version and mask SHA.
"""
from __future__ import annotations

import csv
import hashlib
import json
import logging
import os
import re
import sqlite3
import tempfile
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from .petek_mask_store import PetekProfileStore, private_root
from .petek_surface_v3 import (
    CATEGORY, RECIPE, PetekSafetyError, make_artifact_key,
    require_safe_file, sha256_file,
)

LOG = logging.getLogger("petek.batch")
IMAGE_EXTS = frozenset((".png",".jpg",".jpeg",".webp",".bmp"))
CSV_HEADERS = (
    "Stok Kodu", "Barkod (Dahili)", "Model Kodu", "Kategori",
    "Ham Baskı Dosyası", "Şablon", "Ürün Adı",
    "Mockup Dosya Adı", "Yerel Dosya", "Durum",
)
MAX_ITEMS_PER_RUN = 10000
LOCK_FILENAME = "active_batch.lock"


@dataclass(frozen=True)
class BatchSummary:
    """Stable summary of created, reused, skipped, and rejected images."""
    requested: int
    rendered: int
    reused: int
    rejected: int
    exported: int
    export_csv: str
    export_xlsx: str
    local_only: bool = True
    marketplace_published: bool = False


def safe_code(value: str, label: str) -> str:
    """Limit user-controlled labels to predictable IDs and Excel-safe text."""
    code = str(value or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,48}", code):
        raise PetekSafetyError(f"INVALID_{label}")
    return code


def safe_spreadsheet_text(value: object) -> str:
    """Prevent CSV/Excel formula injection from design file names."""
    text = str(value if value is not None else "")
    return "'" + text if text.lstrip().startswith(("=","+","-","@")) else text


def _safe_scan(base: Path) -> Iterator[Path]:
    """Iterate source designs lazily; never traverse symlink directories."""
    stack=[base]
    while stack:
        folder=stack.pop()
        for entry in sorted(folder.iterdir(),key=lambda p:p.name.casefold()):
            if entry.is_symlink():
                continue
            if entry.is_dir():
                stack.append(entry)
            elif entry.is_file() and entry.suffix.lower() in IMAGE_EXTS:
                yield entry


def safe_design_folder(folder: Path, incoming_root: Path) -> Path:
    """Forbid reading designs outside the user's designated incoming folder."""
    incoming = Path(incoming_root).resolve(strict=True)
    source = Path(folder)
    if source.is_symlink() or not source.is_dir():
        raise PetekSafetyError("DESIGN_DIRECTORY_INVALID")
    resolved = source.resolve(strict=True)
    if resolved != incoming and incoming not in resolved.parents:
        raise PetekSafetyError("DESIGN_DIRECTORY_OUTSIDE_INCOMING")
    return resolved


def _exclusive_file(path: Path, content: bytes) -> None:
    """Publish receipt/report only by non-overwriting same-directory hardlink."""
    path.parent.mkdir(parents=True,exist_ok=True)
    if path.is_symlink():
        raise PetekSafetyError("EXTERNAL_FILE_SYMLINK")
    if path.exists():
        if path.read_bytes()!=content:
            raise PetekSafetyError("EXISTING_REPORT_DIFFERENT")
        return
    fd,tmp=tempfile.mkstemp(dir=str(path.parent),prefix=".petek-report-")
    try:
        with os.fdopen(fd,"wb") as stream:
            stream.write(content);stream.flush();os.fsync(stream.fileno())
        try:
            os.link(tmp,path)
        except FileExistsError:
            if path.read_bytes()!=content:
                raise PetekSafetyError("EXISTING_REPORT_DIFFERENT")
    finally:
        if os.path.exists(tmp):os.unlink(tmp)


class BatchLedger:
    """Isolated SQLite ledger with durable per-artwork/template checkpoints."""

    def __init__(self, root: Path):
        """Create the store in PETEK's private data folder only."""
        self.root=Path(root)
        self.root.mkdir(parents=True,exist_ok=True)
        self.db=self.root/"production.sqlite3"
        if self.db.is_symlink():
            raise PetekSafetyError("UNSAFE_JOB_DB_SYMLINK")
        with self.connect() as db:
            db.execute("""CREATE TABLE IF NOT EXISTS artifacts (
                artifact_key TEXT PRIMARY KEY,
                sku TEXT NOT NULL,
                template_uid TEXT NOT NULL,
                template_sha TEXT NOT NULL,
                mask_sha TEXT NOT NULL,
                art_sha TEXT NOT NULL,
                artwork_name TEXT NOT NULL,
                output_path TEXT NOT NULL,
                output_sha TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL,
                error TEXT NOT NULL DEFAULT '',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            )""")
            db.execute("CREATE INDEX IF NOT EXISTS idx_petek_status ON artifacts(status)")

    def connect(self):
        """Connect with transactions and bounded DB lock waiting."""
        conn=sqlite3.connect(self.db,timeout=30)
        conn.row_factory=sqlite3.Row
        conn.execute("PRAGMA busy_timeout=30000")
        return conn

    def get(self,key:str) -> dict|None:
        """Find a prior immutable production receipt."""
        with self.connect() as db:
            item=db.execute("SELECT * FROM artifacts WHERE artifact_key=?",(key,)).fetchone()
        return dict(item) if item else None

    def reserve(self,fields:dict) -> dict:
        """Reserve immutable source→SKU→output mapping transactionally."""
        with self.connect() as db:
            db.execute("""INSERT OR IGNORE INTO artifacts(
                artifact_key,sku,template_uid,template_sha,mask_sha,art_sha,
                artwork_name,output_path,status)
                VALUES(:artifact_key,:sku,:template_uid,:template_sha,:mask_sha,
                       :art_sha,:artwork_name,:output_path,'PENDING')""",fields)
        result=self.get(fields["artifact_key"])
        for item in ("sku","template_uid","template_sha","mask_sha",
                     "art_sha","artwork_name","output_path"):
            if result[item] != fields[item]:
                raise PetekSafetyError("IMMUTABLE_ARTIFACT_LEDGER_COLLISION")
        return result

    def rendered(self,key:str,output_sha:str) -> None:
        """Atomically checkpoint the completed local image, not commercial publication."""
        with self.connect() as db:
            res=db.execute("""UPDATE artifacts SET output_sha=?,status='RENDERED',
              updated_at=CURRENT_TIMESTAMP
              WHERE artifact_key=? AND status='PENDING'""",(output_sha,key))
            if res.rowcount!=1:
                raise PetekSafetyError("ARTIFACT_STATUS_TRANSITION_REJECTED")

    def mark_uncertain(self,key:str,reason:str) -> None:
        """Quarantine an uncertain artifact; never overwrite existing files."""
        with self.connect() as db:
            db.execute("""UPDATE artifacts SET status='REVIEW_REQUIRED',error=?,
              updated_at=CURRENT_TIMESTAMP WHERE artifact_key=?""",(reason[:400],key))

    def iter_verified(self) -> Iterator[dict]:
        """Read ledger rows incrementally for memory-safe commerce export."""
        db=self.connect()
        try:
            cur=db.execute("""SELECT * FROM artifacts
                WHERE status='RENDERED' ORDER BY sku,template_uid""")
            while True:
                rows=cur.fetchmany(256)
                if not rows:break
                for row in rows:
                    record=dict(row)
                    output=Path(record["output_path"])
                    if (not output.is_file() or output.is_symlink()
                         or sha256_file(output)!=record["output_sha"]):
                        raise PetekSafetyError(
                            "EXPORTED_IMAGE_MISSING_OR_CHANGED:"+record["artifact_key"])
                    yield record
        finally:
            db.close()


class BatchFileLock:
    """Cross-process exclusive batch lock; never auto-break a suspicious lock."""

    def __init__(self,root:Path):
        """Bind lock to PETEK-private runtime root."""
        self.path=Path(root)/LOCK_FILENAME
        self.held=False

    def __enter__(self):
        """Acquire lock with atomic O_EXCL; stale locks demand human inspection."""
        self.path.parent.mkdir(parents=True,exist_ok=True)
        fd=os.open(str(self.path),os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
        with os.fdopen(fd,"w",encoding="utf-8") as stream:
            stream.write(json.dumps({"pid":os.getpid(),"mode":"PETEK_LOCAL_ONLY"}))
            stream.flush();os.fsync(stream.fileno())
        self.held=True
        return self

    def __exit__(self,exc_type,exc_value,traceback):
        """Remove only the lock obtained by this process."""
        if self.held:
            try:self.path.unlink()
            except OSError:LOG.exception("PETEK lock cleanup requires manual review")
        self.held=False


def _row(rec: dict) -> list[str]:
    """Construct a deterministic metadata-only row; barcode is not GS1/EAN."""
    code=rec["sku"]
    return [safe_spreadsheet_text(x) for x in (
        code,"INT-"+code,"M-"+code,CATEGORY,rec["artwork_name"],
        rec["template_uid"],"Dekoratif Ahşap Petek Dolabı — "+
        Path(rec["artwork_name"]).stem,Path(rec["output_path"]).name,
        rec["output_path"],"LOCAL_RENDERED",
    )]


class CommerceExport:
    """Idempotent local CSV/XLSX export without modifying pre-existing files."""

    def __init__(self,export_dir:Path):
        """Bind exports to PETEK-private local directory."""
        self.root=Path(export_dir)
        self.root.mkdir(parents=True,exist_ok=True)

    def write(self,ledger:BatchLedger)->tuple[Path,Path,int]:
        """Stream an immutable CSV and XLSX with bounded memory, verify by re-reading."""
        from itertools import zip_longest
        from openpyxl import Workbook, load_workbook

        receipt=hashlib.sha256()
        count=0
        for record in ledger.iter_verified():
            receipt.update((record["artifact_key"]+":"+record["output_sha"]+"\n").encode())
            count+=1
        digest=receipt.hexdigest()[:24]
        csv_path=self.root/f"PETEK_LOCAL_{digest}.csv"
        xlsx_path=self.root/f"PETEK_LOCAL_{digest}.xlsx"

        if not csv_path.exists():
            fd,temp=tempfile.mkstemp(dir=str(self.root),prefix=".export-",suffix=".csv")
            try:
                with os.fdopen(fd,"w",encoding="utf-8-sig",newline="") as stream:
                    writer=csv.writer(stream)
                    writer.writerow(CSV_HEADERS)
                    for record in ledger.iter_verified():
                        writer.writerow(_row(record))
                    stream.flush()
                    os.fsync(stream.fileno())
                try:
                    os.link(temp,csv_path)
                except FileExistsError:
                    pass
            finally:
                if os.path.exists(temp):
                    os.unlink(temp)

        # Re-validate CSV rather than silently trusting pre-existing files.
        with csv_path.open(encoding="utf-8-sig",newline="") as stream:
            reader=csv.reader(stream)
            if next(reader,None)!=list(CSV_HEADERS):
                raise PetekSafetyError("CSV_EXPORT_HEADER_MISMATCH")
            if any(a!=b for a,b in zip_longest(reader,
                    (_row(record) for record in ledger.iter_verified()))):
                raise PetekSafetyError("CSV_EXPORT_CONTENT_MISMATCH")

        if not xlsx_path.exists():
            workbook=Workbook(write_only=True)
            sheet=workbook.create_sheet("PETEK_YEREL_KAYIT")
            sheet.append(list(CSV_HEADERS))
            for record in ledger.iter_verified():
                sheet.append(_row(record))
            fd,temp=tempfile.mkstemp(dir=str(self.root),prefix=".export-",suffix=".xlsx")
            os.close(fd)
            try:
                workbook.save(temp)
                try:
                    os.link(temp,xlsx_path)
                except FileExistsError:
                    pass
            finally:
                if os.path.exists(temp):
                    os.unlink(temp)

        workbook=load_workbook(xlsx_path,read_only=True,data_only=True)
        try:
            values=workbook.active.iter_rows(values_only=True)
            if list(next(values,()))!=list(CSV_HEADERS):
                raise PetekSafetyError("XLSX_EXPORT_HEADER_MISMATCH")
            if any(list(actual or [])!=expected for actual,expected in zip_longest(
                  values,(_row(record) for record in ledger.iter_verified()))):
                raise PetekSafetyError("XLSX_EXPORT_CONTENT_MISMATCH")
        finally:
            workbook.close()
        return csv_path,xlsx_path,count


class PetekBatchRunner:
    """Streaming, resumable PETEK production for approved masks only."""

    def __init__(self,runtime_root:Path):
        """Use a separate store; never mutate legacy factory DB or original assets."""
        self.runtime_root=Path(runtime_root).resolve(strict=True)
        self.incoming_root=self.runtime_root/"1_GELEN_KAGGLE_BASKILARI"
        self.template_root=self.runtime_root/"2_SABLON_HAVUZU"
        self.store=PetekProfileStore(self.runtime_root,self.template_root)
        self.private=private_root(self.runtime_root)
        self.ledger=BatchLedger(self.private)
        self.cancel=threading.Event()
        self._status={"running":False,"requested":0,"rendered":0,
                      "reused":0,"rejected":0,"last_error":""}

    def status(self)->dict:
        """Return current in-process counter snapshot for UI polling."""
        return dict(self._status)

    def stop(self)->None:
        """Request safe stop between items; never interrupt an atomic publication."""
        self.cancel.set()

    def run(self,design_folder:Path, template_names:list[str],
            store_group:str="MAIN", max_items:int=MAX_ITEMS_PER_RUN)->BatchSummary:
        """Execute approved-only work incrementally with durable resume receipts."""
        group=safe_code(store_group,"STORE_GROUP")
        if not isinstance(template_names,list) or not template_names:
            raise PetekSafetyError("NO_SELECTED_PETEK_TEMPLATES")
        if not 1 <= int(max_items) <= MAX_ITEMS_PER_RUN:
            raise PetekSafetyError("UNSAFE_MAX_ITEMS")
        designs=safe_design_folder(design_folder,self.incoming_root)
        approved=[]
        for filename in sorted(set(template_names)):
            if Path(filename).name!=filename or Path(filename).suffix.lower() not in IMAGE_EXTS:
                raise PetekSafetyError("INVALID_SELECTED_TEMPLATE")
            source=require_safe_file(self.template_root/CATEGORY/filename,
                                     self.template_root/CATEGORY)
            uid=source.stem
            profile,_=self.store.find(uid,source,approved_only=True)
            approved.append((source,uid,profile))
        # Do not erase an early stop request racing the worker startup.
        self._status={"running":True,"requested":0,"rendered":0,
                      "reused":0,"rejected":0,"last_error":""}
        try:
            with BatchFileLock(self.private):
                for art in _safe_scan(designs):
                    if self.cancel.is_set() or self._status["requested"]>=max_items:
                        break
                    art=require_safe_file(art,self.incoming_root)
                    art_sha=sha256_file(art)
                    for src,uid,profile in approved:
                        if self.cancel.is_set() or self._status["requested"]>=max_items:
                            break
                        self._status["requested"]+=1
                        key=make_artifact_key(profile["source_sha256"],art_sha,
                                              profile["mask_sha256"],uid,
                                              front_quad=profile["front_quad"],
                                              physical_aspect=profile["physical_aspect"])
                        # Store group participates in identity: separate sellers
                        # have separate immutable product records even if same art.
                        render_key=key
                        key=hashlib.sha256((key+"|"+group).encode()).hexdigest()[:24]
                        sku="PK-"+group+"-"+key.upper()
                        output=self.private/"outputs"/CATEGORY/(sku+"_"+key+".png")
                        receipt_path=output.with_suffix(".receipt.json")
                        fields=dict(artifact_key=key,sku=sku,template_uid=uid,
                                    template_sha=profile["source_sha256"],
                                    mask_sha=profile["mask_sha256"],art_sha=art_sha,
                                    artwork_name=art.name,output_path=str(output))
                        try:
                            saved=self.ledger.reserve(fields)
                            if saved["status"]=="RENDERED":
                                if not output.is_file() or sha256_file(output)!=saved["output_sha"]:
                                    raise PetekSafetyError("REUSED_OUTPUT_CHANGED_OR_MISSING")
                                receipt=json.loads(receipt_path.read_text(encoding="utf-8"))
                                if (receipt.get("artifact_key") != render_key or
                                    receipt.get("job_key") != key or
                                    receipt.get("output_sha256") != saved["output_sha"]):
                                    raise PetekSafetyError("REUSE_RECEIPT_MISMATCH")
                                self._status["reused"]+=1
                                continue
                            if saved["status"]!="PENDING" or output.exists():
                                raise PetekSafetyError("UNCERTAIN_OUTPUT_REQUIRES_REVIEW")
                            result=self.store.render(uid,src,art,output,marketplace=True)
                            if result["artifact_key"] != render_key:
                                raise PetekSafetyError("RENDER_IDENTITY_CHANGED_DURING_JOB")
                            if sha256_file(output) != result["output_sha256"]:
                                raise PetekSafetyError("OUTPUT_CHANGED_BEFORE_CHECKPOINT")
                            receipt=dict(result, job_key=key, sku=sku)
                            with receipt_path.open("x", encoding="utf-8") as stream:
                                json.dump(receipt, stream, sort_keys=True, allow_nan=False)
                                stream.flush()
                                os.fsync(stream.fileno())
                            self.ledger.rendered(key,result["output_sha256"])
                            self._status["rendered"]+=1
                        except (PetekSafetyError,ValueError,OSError) as exc:
                            self.ledger.mark_uncertain(key,str(exc))
                            self._status["rejected"]+=1
                            self._status["last_error"]=str(exc)
                            LOG.warning("PETEK batch rejected %s: %s",key,exc)
                csv_path,xlsx_path,exported=CommerceExport(
                    self.private/"exports").write(self.ledger)
                return BatchSummary(
                    requested=self._status["requested"],
                    rendered=self._status["rendered"],
                    reused=self._status["reused"],
                    rejected=self._status["rejected"],
                    exported=exported,export_csv=str(csv_path),export_xlsx=str(xlsx_path),
                )
        finally:
            self._status["running"]=False
