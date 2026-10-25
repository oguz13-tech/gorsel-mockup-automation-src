"""Verified, create-only Excel hyperlink helper shared by PETEK and seven-category export."""
from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook

LINK_COL = 9   # Yerel Dosya, 1-based


def xlsx_with_photo_links(headers, data_rows, *, sheet_name):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.freeze_panes = 'A2'
    ws.append(headers)
    for row in data_rows:
        row = list(row)
        path = Path(str(row[LINK_COL - 1]))
        if path.is_symlink() or not path.is_file():
            raise ValueError('Excel için gerçek çıktı dosyası yok/güvensiz: ' + str(path))
        uri = path.resolve(strict=True).as_uri()
        if not uri.startswith('file:///'):
            raise ValueError('Geçersiz yerel görsel bağlantısı: ' + uri)
        ws.append(row)
        cell = ws.cell(row=ws.max_row, column=LINK_COL)
        cell.hyperlink = uri
        cell.style = 'Hyperlink'
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions['A'].width = 27
    ws.column_dimensions['I'].width = 65
    return wb


def verify_photo_links(ws, outputs):
    """Read from normal-mode workbook: ensure row alignment, URI and file existence."""
    for rownum, item in enumerate(outputs, 2):
        c = ws.cell(row=rownum, column=LINK_COL)
        path = Path(item['output'])
        expected = path.resolve(strict=True).as_uri()
        if c.hyperlink is None or c.hyperlink.target != expected or c.value != str(path):
            raise ValueError(f'Excel görsel bağlantısı satır {rownum}: SKU/görsel eşleşmiyor')
