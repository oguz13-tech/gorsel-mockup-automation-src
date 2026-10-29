"""Store-specific product-upload XLSX, built from the EXISTING verified HTTPS
links and production ledger. No second spreadsheet/report and no image scanning.

Each unique print design/category/store is ONE row, even with many independent
template variants. Only user-approved SKU/barcode/model codes are populated.
The operator's actual *blank* marketplace template remains unchanged.
"""
from __future__ import annotations
from collections import defaultdict
import hashlib
import json
import os
from pathlib import Path
import re
import sqlite3
import tempfile
from urllib.parse import urlparse
from openpyxl import load_workbook
from .store_identity_v485 import store_codes,design_number
from .marketplace_excel_v45 import _sha,detect_columns,_normalize
from .cloud_excel_v43 import _valid_https_image_url


def _image_columns(headers):
    # Preserve marketplace field names; recognize numbered additional image URLs.
    result={}
    for i,title in enumerate(headers,1):
        normalized=_normalize(title)
        m=re.fullmatch(r'(?:gorsel|gorsel linki|resim linki|resim|urun resmi|image|image url) ?([1-9][0-9]?)',normalized)
        if m:
            n=int(m.group(1))
            if n in result:raise ValueError('MAGAZA_EXCEL_GORSEL_SUTUNU_TEKRARLI')
            result[n]=i
    return result


def export_store_upload(template_path:Path,state_root:Path,*,store_id:str,report_root:Path,
                        sheet_name:str|None=None)->dict:
    from .store_profiles_v47 import STORE_ID
    if not STORE_ID.fullmatch(store_id) or store_id=='LOCAL':
        raise ValueError('MAGAZA_EXCEL_KIMLIGI_GUVENSIZ')
    template=Path(template_path)
    state=Path(state_root).absolute()
    target_dir=Path(report_root).absolute()
    if not template.is_file() or template.is_symlink() or template.suffix.lower()!='.xlsx' or template.stat().st_size>35_000_000:
        raise ValueError('MAGAZA_EXCEL_SABLONU_XLSX_GUVENSIZ')
    if target_dir.is_symlink() or state not in target_dir.parents:
        raise ValueError('MAGAZA_EXCEL_CIKTI_KLASORU_GUVENSIZ')
    ledger=state/'v4_ortak_uretim.sqlite3'
    cloud=state/'cloud_urls_v43.sqlite3'
    if not ledger.is_file() or not cloud.is_file() or ledger.is_symlink() or cloud.is_symlink():
        raise ValueError('MAGAZA_EXCEL_ONCE_KAYIT_VE_LINK_GEREKLI')
    rows=defaultdict(list)
    root=(state/'uretilen_gorseller').resolve(strict=False)
    with sqlite3.connect(ledger.resolve().as_uri()+'?mode=ro',uri=True) as conn,\
         sqlite3.connect(cloud.resolve().as_uri()+'?mode=ro',uri=True) as urls:
        conn.row_factory=sqlite3.Row
        urls.row_factory=sqlite3.Row
        links={x['artifact_key']:dict(x) for x in urls.execute('SELECT * FROM cloud_link')}
        for rec in conn.execute("SELECT * FROM artifact WHERE status='RENDERED' ORDER BY updated, key"):
            item=dict(rec)
            path=Path(item['output'])
            if path.is_symlink() or root not in path.resolve(strict=False).parents:
                continue
            rel=path.resolve(strict=False).relative_to(root)
            if len(rel.parts)<3 or rel.parts[0]!=store_id:
                continue
            code=store_codes(design_number(item['art_name']),store_id,item['category'])
            if code['sku']!=item['sku']:
                # Never mix old hashed-SKU artifacts into new public code rows.
                continue
            link=links.get(item['key'])
            if not link or link['output_sha']!=item['output_sha'] or not _valid_https_image_url(link['url']):
                continue
            if not path.is_file() or _sha(path)!=item['output_sha']:
                continue
            rows[(code['sku'],item['category'])].append((item['key'],link['url'],code))
    if not rows:
        raise ValueError('MAGAZA_EXCEL_LINKLI_DOGRULANMIS_GORSEL_YOK: '+store_id)
    wb=load_workbook(template)
    try:
        if sheet_name and sheet_name not in wb.sheetnames:
            raise ValueError('MAGAZA_EXCEL_SAYFA_YOK')
        ws=wb[sheet_name] if sheet_name else wb.active
        found=None
        for header_row in range(1,min(ws.max_row,20)+1):
            headers=[ws.cell(header_row,i).value for i in range(1,ws.max_column+1)]
            try:
                cols=detect_columns(headers)
            except ValueError:continue
            if not {'sku','barcode','model'}.issubset(cols):
                raise ValueError('MAGAZA_EXCEL_STOK_BARKOD_MODEL_SUTUNLARI_ZORUNLU')
            image_cols=_image_columns(headers)
            image_cols.setdefault(1,cols['image_url'])
            if len(set(image_cols.values()))!=len(image_cols):
                raise ValueError('MAGAZA_EXCEL_GORSEL_SUTUNU_CAKISMASI')
            found=(header_row,cols,image_cols)
            break
        if found is None:
            raise ValueError('MAGAZA_EXCEL_SABLONUNDA_STOK_BARKOD_MODEL_VE_GORSEL1_BULUNMADI')
        header_row,cols,image_cols=found
        for existing in ws.iter_rows(min_row=header_row+1):
            if any(cell.value is not None for cell in existing):
                raise ValueError('MAGAZA_EXCEL_DOLU_SABLON_KORUNDU')
        unique_rows=sorted(rows.items())
        for n,((sku,category),records) in enumerate(unique_rows,header_row+1):
            codes=records[0][2]
            values={'sku':sku,'barcode':codes['barcode'],'model':codes['model'],
                    'category':category}
            for field in ('sku','barcode','model','category'):
                if field in cols:
                    v=str(values[field])
                    ws.cell(n,cols[field],"'"+v if v.lstrip().startswith(('=','+','-','@')) else v)
            unique_links=list(dict.fromkeys(url for _,url,_ in records))
            if len(unique_links)>len(image_cols):
                raise ValueError('MAGAZA_EXCEL_GORSEL_SUTUNLARI_YETERSIZ: '+sku+
                                 ' / '+str(len(unique_links))+' görsel')
            for j,url in enumerate(unique_links):
                if (j+1) not in image_cols:
                    raise ValueError('MAGAZA_EXCEL_GORSEL_SUTUNU_EKSIK: '+str(j+1))
                cell=ws.cell(n,image_cols[j+1],url)
                cell.hyperlink=url
                cell.style='Hyperlink'
        signature=hashlib.sha256((
            _sha(template)+store_id+json.dumps(
               [(key,[(k,url) for k,url,_ in records]) for key,records in unique_rows],
               ensure_ascii=False,sort_keys=True)).encode()).hexdigest()[:18]
        target_dir.mkdir(parents=True,exist_ok=True)
        path=target_dir/f'{store_id}_TRENDYOL_URUN_YUKLEME_{signature}.xlsx'
        if path.is_symlink():raise ValueError('MAGAZA_EXCEL_CIKTI_SYMLINK')
        if not path.exists():
            fd,tmp=tempfile.mkstemp(prefix='store-xlsx-',suffix='.xlsx',dir=target_dir)
            os.close(fd)
            try:
                wb.save(tmp)
                try:os.link(tmp,path)
                except FileExistsError:pass
            finally:Path(tmp).unlink(missing_ok=True)
    finally:wb.close()
    verified=load_workbook(path,read_only=False)
    try:
        sheet=verified[sheet_name] if sheet_name else verified.active
        for n,((sku,_),records) in enumerate(unique_rows,header_row+1):
            if sheet.cell(n,cols['sku']).value!=sku or sheet.cell(n,cols['barcode']).value!=sku:
                raise ValueError('MAGAZA_EXCEL_KOD_KONTROL_HATASI')
            if sheet.cell(n,cols['model']).value!=records[0][2]['model']:
                raise ValueError('MAGAZA_EXCEL_MODEL_KONTROL_HATASI')
            for j,url in enumerate(dict.fromkeys(u for _,u,_ in records),1):
                cell=sheet.cell(n,image_cols[j])
                if cell.value!=url or not cell.hyperlink or cell.hyperlink.target!=url:
                    raise ValueError('MAGAZA_EXCEL_LINK_KONTROL_HATASI')
    finally:verified.close()
    return {'store':store_id,'output':str(path),'rows':len(unique_rows),
            'images':sum(len(dict.fromkeys(u for _,u,_ in rec)) for _,rec in unique_rows),
            'status':'STORE_TEMPLATE_READY_REVIEW_MARKETPLACE_REQUIREMENTS'}
