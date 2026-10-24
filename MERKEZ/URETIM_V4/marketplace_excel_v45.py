"""Opt-in SKU/HTTPS URL export using the OPERATOR'S actual marketplace template.

Do not invent Trendyol/Hepsiburada schema versions or product barcodes. Detect
common headings only; unsupported headers remain empty for operator completion.
The supplied workbook is never modified. No image upload or marketplace API.
"""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import re
import sqlite3
import tempfile
import unicodedata

from openpyxl import load_workbook


HEADING_ALIASES = {
    'sku': ('stok kodu', 'satici stok kodu', 'satıcı stok kodu', 'seller stock code', 'stock code', 'sku', 'satici sku'),
    'image_url': ('resim linki', 'resim linki 1', 'gorsel 1', 'görsel 1', 'gorsel linki', 'görsel linki',
                  'gorsel linki 1', 'görsel linki 1', 'ana gorsel', 'ana görsel', 'image url', 'image 1', 'imageurl1',
                  'urun gorseli', 'ürün görseli', 'urun resmi 1', 'ürün resmi 1'),
    'category': ('kategori', 'category', 'kategori kodu'),
    'product_name': ('urun adi', 'ürün adı', 'product name', 'urun basligi', 'ürün başlığı', 'baslik', 'başlık'),
    'description': ('aciklama', 'açıklama', 'urun aciklamasi', 'ürün açıklaması', 'description'),
    'barcode': ('barkod', 'barcode', 'ean'),
    'model': ('model kodu', 'model code'),
}


def _normalize(value: object) -> str:
    text = str(value or '').replace('ı', 'i').replace('İ', 'I').casefold()
    text = ''.join(c for c in unicodedata.normalize('NFKD', text) if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9]+', ' ', text)).strip()


def _sha(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for part in iter(lambda: f.read(1024 * 1024), b''):
            h.update(part)
    return h.hexdigest()


def detect_columns(headers: list[object], *, custom: dict[str, str] | None = None) -> dict[str, int]:
    """Explicit custom header-name mappings have precedence over aliases."""
    named = {_normalize(v): n+1 for n, v in enumerate(headers) if str(v or '').strip()}
    columns = {}
    for field, aliases in HEADING_ALIASES.items():
        override = (custom or {}).get(field)
        if override:
            norm = _normalize(override)
            if norm not in named:
                raise ValueError(f'Excel eşleme alanı bulunamadı: {field}: {override}')
            columns[field] = named[norm]
            continue
        for alias in aliases:
            if _normalize(alias) in named:
                columns[field] = named[_normalize(alias)]
                break
    if 'sku' not in columns or 'image_url' not in columns:
        raise ValueError('Excel şablonunda Stok Kodu/SKU ve Görsel 1/Resim Linki sütunları bulunmalı; '
                         'sütun başlıklarını veya eşleme JSON dosyasını kontrol edin.')
    if columns['sku'] == columns['image_url']:
        raise ValueError('SKU ve görsel bağlantısı aynı Excel sütunu olamaz.')
    return columns


def export_from_template(template_path: Path, state_root: Path, rows: list[dict], *,
                         report_root: Path, column_mapping: dict[str, str] | None = None,
                         sheet_name: str | None = None) -> dict:
    """Create an independent, SHA-identifiable .xlsx from verified cloud link DB.

    Existing workbook cells, data, links and formulae are never overwritten;
    only a header-only sheet is populated. Unmapped fields stay intentionally
    blank, because the user/marketplace is the authority for their values.
    """
    template_path = Path(template_path)
    state_root = Path(state_root).absolute()
    report_root = Path(report_root).absolute()
    if template_path.suffix.lower() != '.xlsx' or template_path.is_symlink() or not template_path.is_file():
        raise ValueError('Gerçek bir XLSX Excel şablonu seçin (symlink/XLSM kabul edilmez).')
    if template_path.stat().st_size > 35_000_000:
        raise ValueError('Excel şablonu çok büyük; ayrı boş ürün yükleme XLSX şablonunu seçin.')
    ledger = state_root / 'v4_ortak_uretim.sqlite3'
    cloud = state_root / 'cloud_urls_v43.sqlite3'
    if not ledger.is_file() or not cloud.is_file() or ledger.is_symlink() or cloud.is_symlink():
        raise ValueError('Yerel üretim kaydı / doğrulanmış HTTPS bulut bağlantısı bulunamadı.')
    workbook = load_workbook(template_path)
    try:
        if sheet_name and sheet_name not in workbook.sheetnames:
            raise ValueError('Excel sayfa adı şablonda bulunamadı.')
        ws = workbook[sheet_name] if sheet_name else workbook.active
        # Marketplaces often place instructions and a brand banner above headers.
        # Detect the actual header (first 20 rows) without changing title rows.
        columns = None
        header_row = None
        for candidate in range(1, min(ws.max_row, 20) + 1):
            heads = [ws.cell(candidate, i).value for i in range(1, ws.max_column+1)]
            try:
                columns = detect_columns(heads, custom=column_mapping)
                header_row = candidate
                break
            except ValueError:
                continue
        if columns is None or header_row is None:
            raise ValueError('İlk 20 satırda SKU ve Görsel 1/Resim Linki başlıkları bulunamadı.')
        # Guard against inserting into already-populated products or sample data.
        for row in ws.iter_rows(min_row=header_row+1):
            if any(cell.value is not None for cell in row):
                raise ValueError('Excel sayfasında mevcut ürün/veri var: var olan satırlara yazılmadı. '
                                 'Boş şablon kullanın.')
        with sqlite3.connect(ledger.resolve(strict=True).as_uri()+'?mode=ro', uri=True) as lc, \
             sqlite3.connect(cloud.resolve(strict=True).as_uri()+'?mode=ro', uri=True) as cc:
            lc.row_factory = sqlite3.Row
            cc.row_factory = sqlite3.Row
            db_rows = {r['key']: dict(r) for r in lc.execute('SELECT * FROM artifact WHERE status=\'RENDERED\'')}
            links = {r['artifact_key']: dict(r) for r in cc.execute('SELECT * FROM cloud_link')}
        import urllib.parse
        values = []
        for record in rows:
            key = record.get('key')
            source = db_rows.get(key)
            link = links.get(key)
            if not source or not link or source.get('output_sha') != link.get('output_sha'):
                continue
            url = link['url']
            parsed = urllib.parse.urlparse(url)
            if parsed.scheme != 'https' or not parsed.hostname or parsed.username or parsed.password:
                continue
            output = Path(source['output'])
            # No export of broken/mutated outputs.
            if output.is_symlink() or (state_root/'uretilen_gorseller').resolve(strict=False) not in output.resolve(strict=False).parents or not output.is_file() or _sha(output) != source['output_sha']:
                continue
            sku = str(source['sku'])
            values.append({'sku': sku, 'image_url': url, 'category': source['category'],
                           'product_name': '', 'description': '', 'barcode': '', 'model': ''})
        if not values:
            raise ValueError('SHA doğrulanmış HTTPS bağlantılı üretim bulunamadı; önce Bulut Excel linklemesi yapın.')
        for i, value in enumerate(values, start=header_row+1):
            for field, col in columns.items():
                data = value.get(field, '')
                # Never auto-invent barcodes, SEO content, model identifiers.
                if isinstance(data, str) and data.lstrip().startswith(('=', '+', '-', '@')):
                    data = "'" + data
                cell = ws.cell(i, col, data)
                if field == 'image_url':
                    cell.hyperlink = value[field]
                    cell.style = 'Hyperlink'
        output_hash = hashlib.sha256((_sha(template_path) + '\n' + '\n'.join(v['sku']+v['image_url'] for v in values)
                                      + '\n' + json.dumps(columns, sort_keys=True)).encode()).hexdigest()[:20]
        if report_root.is_symlink() or state_root not in report_root.parents:
            raise ValueError('Rapor klasörü üretim durum dizininde olmalı ve güvenli olmalı.')
        report_root.mkdir(parents=True, exist_ok=True)
        target = report_root / f'PAZARYERI_SABLONUNDAN_{output_hash}.xlsx'
        if not target.exists() and not target.is_symlink():
            fd, temp = tempfile.mkstemp(prefix='excel_v45_', suffix='.xlsx', dir=report_root)
            os.close(fd)
            try:
                workbook.save(temp)
                try:
                    os.link(temp, target)  # create-only; concurrent writers cannot overwrite
                except FileExistsError:
                    pass
            finally:
                Path(temp).unlink(missing_ok=True)
        if target.is_symlink() or not target.is_file():
            raise ValueError('Var olan Excel çıktısı güvensiz.')
        check = load_workbook(target)
        try:
            sheet = check[sheet_name] if sheet_name else check.active
            for i, v in enumerate(values, start=header_row+1):
                if sheet.cell(i, columns['sku']).value != v['sku'] or \
                        sheet.cell(i, columns['image_url']).value != v['image_url'] or \
                        sheet.cell(i, columns['image_url']).hyperlink is None or \
                        sheet.cell(i, columns['image_url']).hyperlink.target != v['image_url']:
                    raise ValueError('Oluşan Excel hücre URL/SKU doğrulaması başarısız.')
        finally:
            check.close()
        return {'output': str(target), 'rows': len(values), 'columns': columns,
                'missing_business_fields': [k for k in ('product_name','description','barcode','model') if k in columns],
                'header_row': header_row, 'template_sha256': _sha(template_path), 'status': 'VERIFY_PRODUCT_FIELDS_MANUALLY'}
    finally:
        workbook.close()
