"""Explicit opt-in cloud URLs for existing, SHA-verified V4 mockups.

Legacy business contract: successful ImgBB/FreeImage HTTPS image URL -> SKU-matched
Excel 'Resim Linki' column. No marketplace publishing. No renderer invocation.
The existing local Excel export remains unchanged; remote results use new files.
"""
from __future__ import annotations
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import sqlite3
from io import StringIO
import tempfile
from datetime import datetime
import threading
from typing import Callable
from urllib.parse import urlparse

import requests
from openpyxl import Workbook, load_workbook
from PETEK_V3 import engine as petek_engine

REMOTE_HEADERS = [
    "Stok Kodu", "Barkod", "Model Kodu", "Kategori", "Ham Baskı Dosyası",
    "Mağaza Grubu", "Ürün Adı", "Açıklama", "Mockup Dosya Adı",
    "Resim Linki", "Eklenme Zamanı", "Yerel Dosya", "Görsel SHA256", "Bulut Durumu",
]


def _keys(runtime_root: Path | None):
    cfg = {}
    # Legacy source of truth: <runtime>/ayarlar.json contains bulut_ayarlari.
    if runtime_root:
        config = Path(runtime_root) / "ayarlar.json"
        if config.is_file() and not config.is_symlink():
            try:
                cfg = json.loads(config.read_text(encoding="utf-8")).get("bulut_ayarlari", {})
            except (OSError, ValueError, AttributeError):
                cfg = {}
    ikeys = cfg.get("imgbb_keys") or os.getenv("IMGBB_KEYS", "").split(",")
    return ([str(k).strip() for k in ikeys if str(k).strip()],
            str(cfg.get("freeimage_key") or os.getenv("FREEIMAGE_KEY", "")).strip())


def _valid_https_image_url(url: str) -> bool:
    if not isinstance(url, str) or len(url) > 2048:
        return False
    u = urlparse(url)
    if (u.scheme != "https" or not u.hostname or u.username is not None
            or u.password is not None or u.fragment or any(ord(c)<32 for c in url)):
        return False
    import ipaddress
    try:
        addr=ipaddress.ip_address(u.hostname)
    except ValueError:
        if u.hostname.lower() in ('localhost',) or u.hostname.lower().endswith(('.localhost','.local')):
            return False
    else:
        if not addr.is_global:
            return False
    return True


def upload_https(path: Path, imgbb_keys: list[str], freeimage_key: str) -> str | None:
    """Legacy official endpoints, with failover; no keys, response bodies or tokens logged."""
    for key in imgbb_keys:
        try:
            with path.open("rb") as f:
                response = requests.post("https://api.imgbb.com/1/upload",
                                         data={"key": key}, files={"image": f}, timeout=30)
            if response.status_code == 200:
                url = str((response.json().get("data") or {}).get("url") or "")
                if _valid_https_image_url(url):
                    return url
        except (requests.RequestException, OSError, ValueError, AttributeError):
            continue
    if freeimage_key:
        try:
            with path.open("rb") as f:
                response = requests.post("https://freeimage.host/api/1/upload",
                    data={"key": freeimage_key, "action": "upload", "format": "json"},
                    files={"source": f}, timeout=30)
            if response.status_code == 200:
                url = str((response.json().get("image") or {}).get("url") or "")
                if _valid_https_image_url(url):
                    return url
        except (requests.RequestException, OSError, ValueError, AttributeError):
            pass
    return None


def _escape(value):
    v = str(value or "")
    return "'" + v if v.lstrip().startswith(("=", "+", "-", "@")) else v


class CloudExcel:
    def __init__(self, root: Path, reports: Path, runtime_root: Path | None = None):
        self.root = Path(root)
        self.reports = Path(reports)
        self.runtime_root = runtime_root
        self.state = self.root / "cloud_urls_v43.sqlite3"
        self.root.mkdir(parents=True, exist_ok=True)
        self.reports.mkdir(parents=True, exist_ok=True)
        if self.state.is_symlink():
            raise ValueError("Bulut link DB symlink olamaz")
        with self.connect() as db:
            db.execute("""CREATE TABLE IF NOT EXISTS cloud_link (
                artifact_key TEXT PRIMARY KEY,
                output_sha TEXT NOT NULL,
                url TEXT NOT NULL,
                recorded_at TEXT DEFAULT CURRENT_TIMESTAMP
            )""")

    def connect(self):
        c = sqlite3.connect(self.state, timeout=30)
        c.row_factory = sqlite3.Row
        c.execute("PRAGMA busy_timeout=30000")
        return c

    def available(self):
        imgbb, freeimage = _keys(self.runtime_root)
        return bool(imgbb or freeimage)

    def _url_for(self, row):
        with self.connect() as c:
            r = c.execute("SELECT output_sha, url FROM cloud_link WHERE artifact_key=?", (row["key"],)).fetchone()
        if r and r["output_sha"] == row["output_sha"] and _valid_https_image_url(r["url"]):
            return r["url"]
        return None

    def publish_existing(self, successful_rows: list[dict], *, notify: Callable = lambda _x: None,
                         stop: threading.Event | None = None, upload_fn=upload_https):
        """Only explicit user action; read existing rendered artifacts; no marketplace API."""
        keys, freeimage = _keys(self.runtime_root)
        if not (keys or freeimage):
            raise RuntimeError("ImgBB / FreeImage anahtarı bulunamadı. ayarlar.json bulut_ayarlari kontrol edin.")
        lock = self.root / "V43_BULUT_LINKLEME.lock"
        try:
            fd = os.open(str(lock), os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        except FileExistsError as ex:
            raise RuntimeError("Bulut Excel linkleme zaten çalışıyor; mevcut işlem durmadan ikinci yükleme başlatılmaz.") from ex
        counters = {"uploaded": 0, "reused": 0, "failed": 0, "skipped": 0}
        original_count=len(successful_rows)
        distinct={}
        for item in successful_rows:
            k=str(item['key'])
            if k in distinct and distinct[k].get('output_sha')!=item.get('output_sha'):
                raise ValueError('BULUT_DUPLICATE_KEY_SHA_CONFLICT')
            distinct[k]=item
        successful_rows=list(distinct.values())
        counters['skipped']=original_count-len(successful_rows)
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as f:
                f.write(str(os.getpid()))
            for row in successful_rows:
                if stop and stop.is_set():
                    break
                path = Path(row["output"])
                safe_root=(self.root / 'uretilen_gorseller').resolve()
                resolved=path.resolve(strict=False)
                if (path.is_symlink() or safe_root not in resolved.parents
                        or not path.is_file() or petek_engine.digest(path) != row["output_sha"]):
                    counters["failed"] += 1
                    notify(f"SHA/çıktı uyuşmazlığı: {row['sku']}; buluta gönderilmedi.")
                    continue
                if self._url_for(row):
                    counters["reused"] += 1
                    continue
                url = upload_fn(path, keys, freeimage)
                if not url or not _valid_https_image_url(url):
                    counters["failed"] += 1
                    notify(f"Yükleme başarısız: {row['sku']}; Excel'e geçersiz URL yazılmayacak.")
                    continue
                # Preserve successful result before Excel export; interrupted
                # Excel phase can be resumed without another image upload.
                with self.connect() as c:
                    c.execute("""INSERT OR REPLACE INTO cloud_link
                        (artifact_key, output_sha, url) VALUES(?,?,?)""",
                        (row["key"], row["output_sha"], url))
                counters["uploaded"] += 1
                notify(f"HTTPS link oluşturuldu: {row['sku']}")
            csv_out, xlsx_out, count = self.export(successful_rows)
            return {**counters, "linked": count, "xlsx": str(xlsx_out), "csv": str(csv_out),
                    "status": "PARTIAL" if counters["failed"] else ("STOPPED" if stop and stop.is_set() else "LINKS_READY")}
        finally:
            lock.unlink(missing_ok=True)

    def export(self, rows):
        selected = []
        for row in rows:
            url = self._url_for(row)
            if not url:
                continue
            path = Path(row["output"])
            safe_root=(self.root / 'uretilen_gorseller').resolve()
            resolved=path.resolve(strict=False)
            if (path.is_symlink() or safe_root not in resolved.parents
                    or not path.is_file() or petek_engine.digest(path) != row["output_sha"]):
                continue
            cat = row["category"]
            sku = row["sku"]
            group_name = sku.split("-", 1)[1].rsplit("-", 1)[0] if "-" in sku else "LOCAL"
            selected.append([
                _escape(v) for v in (
                    sku, "DAHILI-" + sku, "M-" + sku, cat, row['art_name'],
                    group_name, f"{cat} - {Path(row['art_name']).stem}",
                    "",  # SEO açıklaması üretilmediyse ASLA uydurma
                    path.name, url, datetime.now().strftime("%d.%m.%Y %H:%M:%S"),
                    str(path), row['output_sha'], "BULUT_LINK_HAZIR")
            ])
        fingerprint = hashlib.sha256("|".join(x[0] + x[9] for x in selected).encode()).hexdigest()[:18]
        csv_out = self.reports / f"ORTAK_7_KATEGORI_{fingerprint}_BULUT_LINKLI.csv"
        xlsx_out = self.reports / f"ORTAK_7_KATEGORI_{fingerprint}_BULUT_LINKLI.xlsx"
        if not csv_out.exists():
            buffer = StringIO(newline="")
            w = csv.writer(buffer); w.writerow(REMOTE_HEADERS); w.writerows(selected)
            petek_engine.atomic_new_file(csv_out, ("\ufeff" + buffer.getvalue()).encode("utf-8"))
        if not xlsx_out.exists():
            wb = Workbook(); ws = wb.active; ws.title = "BULUT_LINKLI_7_KATEGORI"
            ws.freeze_panes = "A2"; ws.append(REMOTE_HEADERS)
            for line in selected:
                ws.append(line)
                cell = ws.cell(ws.max_row, 10)
                cell.hyperlink = line[9]; cell.style = "Hyperlink"
                ws.cell(ws.max_row, 12).hyperlink = Path(line[11]).resolve().as_uri()
                ws.cell(ws.max_row, 12).style = "Hyperlink"
            ws.auto_filter.ref = ws.dimensions
            ws.column_dimensions["J"].width = 68
            ws.column_dimensions["L"].width = 64
            fd, tmp = tempfile.mkstemp(prefix="v43excel-", suffix=".xlsx", dir=str(self.reports))
            os.close(fd)
            try:
                wb.save(tmp)
                petek_engine.atomic_new_file(xlsx_out, Path(tmp).read_bytes())
            finally:
                Path(tmp).unlink(missing_ok=True)
        # Read back and confirm every reported SKU has its exact HTTPS URL.
        wb = load_workbook(xlsx_out)
        try:
            ws = wb.active
            assert ws.max_row == len(selected) + 1
            for n, row in enumerate(selected, 2):
                c = ws.cell(n, 10)
                if c.value != row[9] or c.hyperlink is None or c.hyperlink.target != row[9]:
                    raise RuntimeError("Excel URL/SKU doğrulama hatası: " + row[0])
        finally:
            wb.close()
        return csv_out, xlsx_out, len(selected)