"""P0 Regression Tests - Real assertions, not stubs."""
import os
import sys
import json
import glob
import unittest
import py_compile
import tempfile
import sqlite3
import shutil
from pathlib import Path
from unittest.mock import patch

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(BASE_DIR, "99_SISTEM_MOTORLARI_VE_KODLAR"))

try:
    from SISTEM_MOTORU.pc_master_factory import StateManager, AsyncCloudUploader, FastExcelManager
    IMPORT_SUCCESS = True
except ImportError as e:
    IMPORT_SUCCESS = False
    IMPORT_ERROR = str(e)


class TestProductionModuleCompiles(unittest.TestCase):
    """Verify that the production module has no SyntaxErrors."""

    def test_pc_master_factory_compiles(self):
        pc_path = os.path.join(BASE_DIR, "99_SISTEM_MOTORLARI_VE_KODLAR",
                               "SISTEM_MOTORU", "pc_master_factory.py")
        try:
            py_compile.compile(pc_path, doraise=True)
        except py_compile.PyCompileError as e:
            self.fail(f"SyntaxError in pc_master_factory.py: {e}")
            
    def test_module_imports(self):
        self.assertTrue(IMPORT_SUCCESS, f"Could not import pc_master_factory: {globals().get('IMPORT_ERROR')}")

class TestSecretsRemoved(unittest.TestCase):
    """Verify no API keys remain in any ayarlar*.json file."""

    def test_no_secrets_in_any_json(self):
        json_files = glob.glob(os.path.join(BASE_DIR, "**", "ayarlar*.json"), recursive=True)
        self.assertGreater(len(json_files), 0, "No ayarlar JSON files found")

        for jf in json_files:
            with open(jf, "r", encoding="utf-8") as f:
                data = json.load(f)
            cloud = data.get("bulut_ayarlari", {})

            # imgbb_keys must be empty list
            keys = cloud.get("imgbb_keys", [])
            self.assertEqual(keys, [], f"imgbb_keys not empty in {jf}")

            # freeimage_key must be empty string
            fk = cloud.get("freeimage_key", "")
            self.assertEqual(fk, "", f"freeimage_key not empty in {jf}")

            # cloudinary secrets must be empty
            cl = cloud.get("cloudinary", {})
            self.assertEqual(cl.get("api_key", ""), "", f"cloudinary api_key not empty in {jf}")
            self.assertEqual(cl.get("api_secret", ""), "", f"cloudinary api_secret not empty in {jf}")
            
    def test_no_hardcoded_secrets_in_code(self):
        """Scan production python code for hardcoded API keys or passwords."""
        import re
        py_files = glob.glob(os.path.join(BASE_DIR, "99_SISTEM_MOTORLARI_VE_KODLAR", "**", "*.py"), recursive=True)
        
        # We look for common patterns but ignore empty assignments or our own test variables.
        # This regex looks for variables like 'api_key = "some_long_string"'
        secret_pattern = re.compile(r'(?i)(api_key|secret|password|token)\s*=\s*[\'"](?!\s*[\'"])([a-zA-Z0-9_\-]{8,})[\'"]')
        
        for pf in py_files:
            with open(pf, "r", encoding="utf-8") as f:
                content = f.read()
                
            matches = secret_pattern.findall(content)
            self.assertEqual(len(matches), 0, f"Found hardcoded secret patterns in {pf}: {matches}")


class TestSQLiteBootstrap(unittest.TestCase):
    """Verify production module can bootstrap a temp SQLite DB."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.db_path = Path(self.tmp) / "test_state.db"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_old_recipe_failure_does_not_block_current_product_completion(self):
        from SISTEM_MOTORU.pc_master_factory import StateManager

        sm = StateManager(db_path=self.db_path)
        with sm.conn:
            sm.conn.execute(
                """INSERT INTO masters (
                       master_id, theme_id, theme_title, style_key,
                       horizontal_path, vertical_path
                   ) VALUES (?, ?, ?, ?, ?, ?)""",
                ("M00001", 1, "Tema", "style", "H.png", "V.png"),
            )
            sm.conn.execute(
                """INSERT INTO products (
                       sku, master_id, kategori_key, barkod, model_kodu, status
                   ) VALUES (?, ?, ?, ?, ?, ?)""",
                ("sg-M00001", "M00001", "01_Wifi_Kutusu", "sg-M00001", "M-sg-M00001", "UPLOADED_PENDING_EXCEL"),
            )
            sm.conn.execute(
                """INSERT INTO render_artifacts (
                       artifact_key, sku, status
                   ) VALUES (?, ?, ?)""",
                ("old_source_01_Wifi_Kutusu_old_template_v2_old", "sg-M00001", "UPLOAD_FAILED"),
            )
            sm.conn.execute(
                """INSERT INTO render_artifacts (
                       artifact_key, sku, status
                   ) VALUES (?, ?, ?)""",
                ("new_source_01_Wifi_Kutusu_new_template_v3_epic3_20261005_r2", "sg-M00001", "UPLOADED_PENDING_EXCEL"),
            )

        sm.mark_artifacts_completed(
            ["new_source_01_Wifi_Kutusu_new_template_v3_epic3_20261005_r2"],
            recipe_version="v3_epic3_20261005_r2",
        )

        row = sm.conn.execute(
            "SELECT status FROM products WHERE sku=?",
            ("sg-M00001",),
        ).fetchone()
        self.assertEqual(row["status"], "COMPLETED")

        old = sm.conn.execute(
            "SELECT status FROM render_artifacts WHERE artifact_key LIKE 'old_source%'"
        ).fetchone()
        self.assertEqual(
            old["status"],
            "UPLOAD_FAILED",
            "Historical recipe evidence must remain unchanged for audit/recovery",
        )

    def test_state_manager_creates_db(self):
        sm = StateManager(db_path=self.db_path)
        
        # Insert a failed artifact
        with sm.conn:
            sm.conn.execute(
                "INSERT INTO render_artifacts (artifact_key, status) VALUES (?, ?)", 
                ("art-001", "UPLOAD_FAILED")
            )
        
        # Verify UPLOAD_FAILED is NOT considered completed
        is_completed = sm.is_artifact_completed("art-001")
        self.assertFalse(is_completed, "UPLOAD_FAILED must not be treated as completed by StateManager")
        
        # Upload success alone is not end-to-end completion. The Excel row must
        # be durably flushed before the artifact becomes COMPLETED.
        sm.save_render_artifact({
            "artifact_key": "art-001",
            "upload_url": "http://test.url",
            "status": "UPLOADED_PENDING_EXCEL"
        })
        self.assertFalse(
            sm.is_artifact_completed("art-001"),
            "Upload-only state must not be treated as completed",
        )
        sm.mark_artifacts_completed(["art-001"])
        self.assertTrue(
            sm.is_artifact_completed("art-001"),
            "Only COMPLETED proves render + upload + Excel finished",
        )


class TestArtifactCompletionRequiresLocalFile(unittest.TestCase):
    def test_completed_db_record_without_file_is_not_reusable(self):
        from SISTEM_MOTORU.pc_master_factory import artifact_record_is_complete_on_disk

        self.assertFalse(
            artifact_record_is_complete_on_disk({
                "status": "COMPLETED",
                "local_output_path": str(Path(tempfile.gettempdir()) / "missing-render.jpg"),
            })
        )

    def test_completed_db_record_with_file_is_reusable(self):
        from SISTEM_MOTORU.pc_master_factory import artifact_record_is_complete_on_disk

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "render.jpg"
            path.write_bytes(b"ok")
            self.assertTrue(
                artifact_record_is_complete_on_disk({
                    "status": "COMPLETED",
                    "local_output_path": str(path),
                })
            )


class TestForegroundAIDependencyPreflight(unittest.TestCase):
    def test_missing_sam2_symbols_are_reported(self):
        from types import SimpleNamespace
        from SISTEM_MOTORU.pc_master_factory import foreground_ai_dependency_status

        fake_transformers = SimpleNamespace(
            AutoProcessor=object(),
            AutoModelForZeroShotObjectDetection=object(),
        )

        with patch(
            "SISTEM_MOTORU.pc_master_factory.importlib.util.find_spec",
            return_value=object(),
        ), patch(
            "SISTEM_MOTORU.pc_master_factory.importlib.import_module",
            return_value=fake_transformers,
        ):
            status = foreground_ai_dependency_status()

        self.assertTrue(status["torch"])
        self.assertTrue(status["transformers"])
        self.assertFalse(status["transformers_required_symbols"])
        self.assertIn("transformers.Sam2Processor", status["missing"])
        self.assertIn("transformers.Sam2Model", status["missing"])

    def test_complete_transformers_api_passes_symbol_check(self):
        from types import SimpleNamespace
        from SISTEM_MOTORU.pc_master_factory import foreground_ai_dependency_status

        fake_transformers = SimpleNamespace(
            AutoProcessor=object(),
            AutoModelForZeroShotObjectDetection=object(),
            Sam2Processor=object(),
            Sam2Model=object(),
        )

        with patch(
            "SISTEM_MOTORU.pc_master_factory.importlib.util.find_spec",
            return_value=object(),
        ), patch(
            "SISTEM_MOTORU.pc_master_factory.importlib.import_module",
            return_value=fake_transformers,
        ):
            status = foreground_ai_dependency_status()

        self.assertTrue(status["transformers_required_symbols"])
        self.assertEqual(status["missing"], [])


class TestDuplicateProfileIdentityGuard(unittest.TestCase):
    def test_duplicate_profile_uids_returns_only_repeated_nonempty_ids(self):
        from SISTEM_MOTORU.pc_master_factory import duplicate_profile_uids

        rows = [
            {"template_uid": "A"},
            {"template_uid": "B"},
            {"template_uid": "A"},
            {"template_uid": ""},
            {"template_uid": None},
            {"template_uid": "C"},
            {"template_uid": "C"},
        ]
        self.assertEqual(duplicate_profile_uids(rows), ["A", "C"])


class TestForegroundPlanSemantics(unittest.TestCase):
    def test_default_review_plan_is_not_occlusion_evidence(self):
        from SISTEM_MOTORU.pc_master_factory import (
            foreground_plan_has_positive_selection,
            foreground_is_required,
        )

        default_plan = {
            "object_boxes": [],
            "add_brush": [],
            "subtract_brush": [],
            "feather": 1.2,
            "grabcut_iterations": 5,
        }
        self.assertFalse(foreground_plan_has_positive_selection(default_plan))
        self.assertFalse(foreground_is_required({"has_occlusion": False}, default_plan))

    def test_positive_manual_plan_or_profile_requires_foreground(self):
        from SISTEM_MOTORU.pc_master_factory import (
            foreground_plan_has_positive_selection,
            foreground_is_required,
        )

        positive_plan = {
            "object_boxes": [[10, 10, 30, 30]],
            "add_brush": [],
            "subtract_brush": [],
        }
        self.assertTrue(foreground_plan_has_positive_selection(positive_plan))
        self.assertTrue(foreground_is_required({"has_occlusion": False}, positive_plan))
        self.assertTrue(foreground_is_required({"has_occlusion": True}, {}))

    def test_subtract_only_plan_is_not_positive_evidence(self):
        from SISTEM_MOTORU.pc_master_factory import foreground_plan_has_positive_selection

        self.assertFalse(
            foreground_plan_has_positive_selection({
                "object_boxes": [],
                "add_brush": [],
                "subtract_brush": [{"points": [[1, 1], [2, 2]]}],
            })
        )


class TestForegroundCacheTrust(unittest.TestCase):
    def test_required_foreground_rejects_empty_cache(self):
        from SISTEM_MOTORU.pc_master_factory import cached_foreground_mask_usable
        import numpy as np

        self.assertFalse(
            cached_foreground_mask_usable(
                np.zeros((10, 10), dtype=np.uint8),
                True,
            )
        )
        mask = np.zeros((10, 10), dtype=np.uint8)
        mask[3:6, 3:6] = 255
        self.assertTrue(cached_foreground_mask_usable(mask, True))

    def test_optional_foreground_allows_empty_cache(self):
        from SISTEM_MOTORU.pc_master_factory import cached_foreground_mask_usable
        import numpy as np

        self.assertTrue(
            cached_foreground_mask_usable(
                np.zeros((10, 10), dtype=np.uint8),
                False,
            )
        )


class TestRuntimeWorkspaceSeparation(unittest.TestCase):
    def test_runtime_workspace_override_is_explicit_and_resolved(self):
        from SISTEM_MOTORU.pc_master_factory import resolve_runtime_workspace

        with tempfile.TemporaryDirectory() as source_tmp, tempfile.TemporaryDirectory() as runtime_tmp:
            source = Path(source_tmp)
            runtime = Path(runtime_tmp)
            with patch.dict(
                os.environ,
                {"MOCKUP_RUNTIME_WORKSPACE": str(runtime)},
                clear=False,
            ):
                resolved = resolve_runtime_workspace(source)

            self.assertEqual(resolved, runtime.resolve())
            self.assertNotEqual(resolved, source.resolve())

    def test_runtime_workspace_defaults_to_source_for_backward_compatibility(self):
        from SISTEM_MOTORU.pc_master_factory import resolve_runtime_workspace

        with tempfile.TemporaryDirectory() as source_tmp:
            source = Path(source_tmp)
            with patch.dict(
                os.environ,
                {"MOCKUP_RUNTIME_WORKSPACE": ""},
                clear=False,
            ):
                self.assertEqual(
                    resolve_runtime_workspace(source),
                    source.resolve(),
                )


class TestENVFallbackBehavior(unittest.TestCase):
    """Test that AsyncCloudUploader config values correctly fall through to ENV."""

    def test_imgbb_empty_list_filters_blanks(self):
        with patch.dict(os.environ, {"IMGBB_KEYS": ""}):
            with patch("SISTEM_MOTORU.pc_master_factory.AYARLAR", {"bulut_ayarlari": {"imgbb_keys": []}}):
                uploader = AsyncCloudUploader()
                self.assertEqual(uploader.imgbb_keys, [], "Empty env should produce empty list, not ['']")

    def test_imgbb_valid_keys_parsed(self):
        with patch.dict(os.environ, {"IMGBB_KEYS": "key1,key2"}):
            with patch("SISTEM_MOTORU.pc_master_factory.AYARLAR", {"bulut_ayarlari": {"imgbb_keys": []}}):
                uploader = AsyncCloudUploader()
                self.assertEqual(uploader.imgbb_keys, ["key1", "key2"])

    def test_freeimage_fallback(self):
        with patch.dict(os.environ, {"FREEIMAGE_KEY": "env_freeimage"}):
            with patch("SISTEM_MOTORU.pc_master_factory.AYARLAR", {"bulut_ayarlari": {"freeimage_key": ""}}):
                uploader = AsyncCloudUploader()
                self.assertEqual(uploader.freeimage_key, "env_freeimage")


class TestExcelFlushNoDeadlock(unittest.TestCase):
    """Verify real Excel write operations don't deadlock with threading."""

    def setUp(self):
        self.tmp = tempfile.mktemp(suffix=".xlsx")

    def tearDown(self):
        if os.path.exists(self.tmp):
            os.remove(self.tmp)

    def test_50_row_excel_flush(self):
        import threading
        
        manager = FastExcelManager(file_path=Path(self.tmp))
        
        def write_rows(start, count):
            for i in range(start, start+count):
                row_data = {
                    "Stok Kodu": f"SKU-{i:04d}",
                    "Barkod": f"BRK-{i}",
                    "Model Kodu": "MODEL1",
                    "Kategori": "TestKategori",
                    "Ham Baskı Dosyası": "a.jpg",
                    "Mağaza Grubu": "GRUP",
                    "Ürün Adı": "Urun",
                    "Açıklama": "Aciklama",
                    "Mockup Dosya Adı": "m.jpg",
                    "Resim Linki": "http://x",
                    "Eklenme Zamanı": "now"
                }
                manager.append_row(row_data)
                
        threads = []
        # Create 5 threads, each writing 10 rows
        for i in range(5):
            t = threading.Thread(target=write_rows, args=(i*10, 10))
            threads.append(t)
            t.start()
            
        for t in threads:
            t.join(timeout=5.0)
            self.assertFalse(t.is_alive(), "Thread deadlocked during Excel append")
            
        # Verify file has 51 rows (header + 50 data)
        import openpyxl
        wb2 = openpyxl.load_workbook(self.tmp)
        ws2 = wb2.active
        self.assertEqual(ws2.max_row, 51, f"Expected 51 rows, got {ws2.max_row}")
        wb2.close()


    def test_transient_threshold_flush_keeps_rows_for_final_retry(self):
        manager = FastExcelManager(file_path=Path(self.tmp))

        def row_data(i):
            return {
                "Stok Kodu": f"SKU-{i:04d}",
                "Barkod": f"BRK-{i}",
                "Model Kodu": "MODEL1",
                "Kategori": "TestKategori",
                "Ham Baskı Dosyası": "a.jpg",
                "Mağaza Grubu": "GRUP",
                "Ürün Adı": "Urun",
                "Açıklama": "Aciklama",
                "Mockup Dosya Adı": f"m-{i}.jpg",
                "Resim Linki": "http://x",
                "Eklenme Zamanı": "now",
            }

        for i in range(49):
            self.assertTrue(manager.append_row(row_data(i)))

        with patch.object(manager, "flush", return_value=False):
            accepted = manager.append_row(row_data(49))

        self.assertTrue(
            accepted,
            "A row already queued in memory must remain pending even if the "
            "threshold flush is transiently unsuccessful",
        )
        self.assertEqual(len(manager._buffer), 50)

        self.assertTrue(manager.flush(), "Final run-level flush should retry buffer")
        import openpyxl
        wb = openpyxl.load_workbook(self.tmp)
        try:
            self.assertEqual(wb.active.max_row, 51)
        finally:
            wb.close()


class TestDetailPreservingRealism(unittest.TestCase):
    def test_quality_render_scale_supersamples_small_templates(self):
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        scale = MockupRenderer.compute_quality_render_scale(
            1200, 1600, 2400, 3200, max_scale=2.5
        )
        self.assertAlmostEqual(scale, 2.0, places=5)

    def test_quality_render_scale_does_not_downscale_large_templates(self):
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        scale = MockupRenderer.compute_quality_render_scale(
            3000, 4000, 2400, 3200, max_scale=2.5
        )
        self.assertEqual(scale, 1.0)

    def test_microdetail_restoration_preserves_shape_and_dtype(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        base = np.full((80, 80, 3), 100, dtype=np.uint8)
        art = base.copy()
        art[20:60, 39:41] = 240
        mask = np.zeros((80, 80), dtype=np.float32)
        mask[10:70, 10:70] = 1.0

        restored = MockupRenderer.restore_artwork_microdetail(
            base, art, mask, strength=0.5
        )
        self.assertEqual(restored.shape, base.shape)
        self.assertEqual(restored.dtype, np.uint8)
        self.assertGreater(int(restored[40, 40, 0]), int(base[40, 40, 0]))



class TestEdgeSafeMasterQuality(unittest.TestCase):
    def test_nominal_aspect_compatibility_does_not_replicate_edges(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        image = np.zeros((896, 1200, 3), dtype=np.uint8)
        image[:, 0, :] = 77
        image[:, -1, :] = 199
        prepared = MockupRenderer.prepare_artwork_for_nominal_aspect(
            image,
            1.42857143,
        )
        self.assertEqual(prepared.shape, image.shape)
        self.assertTrue(np.array_equal(prepared, image))

    def test_large_aspect_mismatch_is_rejected(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        image = np.zeros((1000, 1000, 3), dtype=np.uint8)
        with self.assertRaises(ValueError):
            MockupRenderer.prepare_artwork_for_nominal_aspect(
                image,
                2.0,
            )

    def test_final_master_quality_preserves_flat_regions(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        image = np.full((160, 160, 3), 128, dtype=np.uint8)
        enhanced = MockupRenderer.final_master_quality_pass(image)
        self.assertEqual(enhanced.shape, image.shape)
        self.assertEqual(enhanced.dtype, np.uint8)
        self.assertLessEqual(
            int(np.abs(enhanced.astype(np.int16) - image.astype(np.int16)).max()),
            1,
        )

    def test_final_master_quality_keeps_values_bounded(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        image = np.zeros((160, 160, 3), dtype=np.uint8)
        image[:, 80:, :] = 255
        enhanced = MockupRenderer.final_master_quality_pass(image)
        self.assertGreaterEqual(int(enhanced.min()), 0)
        self.assertLessEqual(int(enhanced.max()), 255)



class TestR6AdaptiveRealism(unittest.TestCase):
    def test_subtle_glass_streak_is_small_and_low_opacity(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        streak = MockupRenderer.generate_subtle_glass_streak(
            400, 300, intensity=0.075
        )
        self.assertEqual(streak.shape, (300, 400, 3))
        self.assertGreater(float(streak.max()), 0.0)
        self.assertLessEqual(float(streak.max()), 0.12 + 1e-6)
        active_fraction = float(np.mean(streak[:, :, 0] > 0.01))
        self.assertLess(active_fraction, 0.35)

    def test_detail_aware_reflection_preserves_some_reflection(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        art = np.zeros((120, 160, 3), dtype=np.uint8)
        art[:, 80:, :] = 230
        art[20:100:4, 20:75, :] = 255
        reflection = np.full((120, 160, 3), 0.20, dtype=np.float32)
        mask = np.ones((120, 160), dtype=np.float32)

        protected = MockupRenderer.protect_artwork_detail_from_reflection(
            reflection,
            art,
            mask,
            min_reflection_fraction=0.30,
        )
        self.assertEqual(protected.shape, reflection.shape)
        self.assertGreater(float(protected.max()), 0.0)
        self.assertLess(float(np.mean(protected[:, 90:, :])), 0.20)
        self.assertGreaterEqual(float(protected.min()), 0.0)

    def test_smart_marketplace_focus_returns_exact_canvas(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        image = np.full((600, 1000, 3), 110, dtype=np.uint8)
        image[180:430, 230:800, :] = 180
        quad = [[230, 180], [800, 180], [800, 430], [230, 430]]

        result = MockupRenderer.prepare_marketplace_canvas(
            image,
            target_width=600,
            target_height=800,
            product_quad=quad,
            smart_product_focus=True,
        )
        self.assertEqual(result.shape, (800, 600, 3))
        self.assertEqual(result.dtype, np.uint8)



class TestFinalMarketplaceQA(unittest.TestCase):
    def _textured_image(self, width=600, height=800):
        import numpy as np

        y, x = np.mgrid[0:height, 0:width]
        base = ((x * 3 + y * 5) % 180 + 35).astype(np.uint8)
        return np.stack(
            [base, np.roll(base, 3, axis=1), np.roll(base, 5, axis=0)],
            axis=-1,
        )

    def test_final_qa_accepts_valid_marketplace_image(self):
        from SISTEM_MOTORU.pc_master_factory import QualityInspector

        image = self._textured_image()
        quad = [[80, 160], [520, 160], [520, 560], [80, 560]]
        valid, message = QualityInspector.inspect_final_output(
            image,
            expected_width=600,
            expected_height=800,
            product_quad=quad,
            smart_product_focus=True,
            foreground_fill_fraction=0.85,
        )
        self.assertTrue(valid, message)

    def test_final_qa_rejects_product_cropped_outside_canvas(self):
        from SISTEM_MOTORU.pc_master_factory import QualityInspector

        image = self._textured_image()
        quad = [[-20, 120], [520, 120], [520, 560], [-20, 560]]
        valid, message = QualityInspector.inspect_final_output(
            image,
            expected_width=600,
            expected_height=800,
            product_quad=quad,
        )
        self.assertFalse(valid)
        self.assertIn("kırpılmış", message)

    def test_final_qa_rejects_excessive_smart_focus_letterboxing(self):
        from SISTEM_MOTORU.pc_master_factory import QualityInspector

        image = self._textured_image()
        quad = [[80, 160], [520, 160], [520, 560], [80, 560]]
        valid, message = QualityInspector.inspect_final_output(
            image,
            expected_width=600,
            expected_height=800,
            product_quad=quad,
            smart_product_focus=True,
            foreground_fill_fraction=0.40,
        )
        self.assertFalse(valid)
        self.assertIn("yeterince dolduramadı", message)

    def test_smart_canvas_metadata_tracks_product_inside_frame(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        image = np.full((600, 1000, 3), 110, dtype=np.uint8)
        image[180:430, 230:800, :] = 180
        quad = [[230, 180], [800, 180], [800, 430], [230, 430]]

        result, meta = MockupRenderer.prepare_marketplace_canvas(
            image,
            target_width=600,
            target_height=800,
            product_quad=quad,
            smart_product_focus=True,
            return_metadata=True,
        )
        self.assertEqual(result.shape, (800, 600, 3))
        self.assertGreater(meta["foreground_fill_fraction"], 0.55)
        out_quad = np.asarray(meta["product_quad_output"], dtype=np.float32)
        self.assertGreaterEqual(float(out_quad[:, 0].min()), -1.0)
        self.assertGreaterEqual(float(out_quad[:, 1].min()), -1.0)
        self.assertLessEqual(float(out_quad[:, 0].max()), 600.0)
        self.assertLessEqual(float(out_quad[:, 1].max()), 800.0)


class TestPrintMaskAlphaSemantics(unittest.TestCase):
    def test_fully_opaque_alpha_does_not_hide_artwork(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        quad = np.zeros((100, 100), dtype=np.float32)
        quad[10:90, 10:90] = 1.0
        opaque_alpha = np.full((100, 100), 255, dtype=np.uint8)

        resolved = MockupRenderer.resolve_print_mask(
            quad,
            opaque_alpha,
            100,
            100,
        )
        self.assertGreater(np.count_nonzero(resolved > 0.05), 6000)
        self.assertTrue(np.allclose(resolved, quad))

    def test_meaningful_transparent_hole_restricts_quad(self):
        import numpy as np
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer

        quad = np.zeros((100, 100), dtype=np.float32)
        quad[10:90, 10:90] = 1.0
        alpha = np.full((100, 100), 255, dtype=np.uint8)
        alpha[25:75, 25:75] = 0

        resolved = MockupRenderer.resolve_print_mask(
            quad,
            alpha,
            100,
            100,
        )
        self.assertGreater(np.count_nonzero(resolved > 0.05), 2000)
        self.assertLess(np.count_nonzero(resolved > 0.05), np.count_nonzero(quad > 0.05))


class TestQAFailStopsProduction(unittest.TestCase):
    """Verify QA failure returns False (quarantine), not just a warning."""
    
    @patch("SISTEM_MOTORU.pc_master_factory.QualityInspector.inspect_mockup")
    @patch("SISTEM_MOTORU.pc_master_factory.MockupRenderer.render_cv2_optik")
    @patch("SISTEM_MOTORU.pc_master_factory.np.fromfile")
    @patch("SISTEM_MOTORU.pc_master_factory.cv2.imdecode")
    @patch("SISTEM_MOTORU.pc_master_factory.cv2.imencode")
    def test_qa_fail_returns_false(self, mock_imencode, mock_imdecode, mock_fromfile, mock_render, mock_qa):
        from SISTEM_MOTORU.pc_master_factory import MockupRenderer
        import numpy as np
        import cv2
        
        # Mock dependencies to reach QA logic safely
        mock_fromfile.return_value = np.array([])
        mock_imdecode.return_value = np.zeros((10, 20, 3), dtype=np.uint8)
        mock_render.return_value = np.zeros((10, 10, 3), dtype=np.uint8)
        mock_qa.return_value = (False, "Mocked QA fail")
        mock_imencode.return_value = (True, np.array([1]))
        
        # Create dummy paths
        tmp_dir = Path(tempfile.mkdtemp())
        raw_path = tmp_dir / "fake_raw.jpg"
        template_path = tmp_dir / "fake_template.png"
        output_path = tmp_dir / "fake_output.jpg"
        
        cv2.imwrite(str(raw_path), np.zeros((10,10,3), dtype=np.uint8))
        cv2.imwrite(str(template_path), np.zeros((10,10,3), dtype=np.uint8))
        
        try:
            # We must mock Path.exists dynamically inside process_mockup ONLY for orig_path,
            # or just mock Path.exists everywhere to True and let mock_fromfile handle the non-existent file
            # Wait, if mock_exists is True everywhere, process_mockup calls orig_path.exists() -> True
            # then calls np.fromfile(str(orig_path)). BUT we already mocked np.fromfile!
            # So we don't need the file to exist on disk!
            with patch("SISTEM_MOTORU.pc_master_factory.Path.exists", return_value=True):
                # Run process_mockup
                result = MockupRenderer.process_mockup(
                    raw_path,
                    template_path,
                    {"profile_orientation": "H", "nominal_aspect": 2.0},
                    False,
                    output_path,
                )
                
                # QA should fail and cause process_mockup to return False
                self.assertFalse(result, "process_mockup should return False on QA failure")
                mock_qa.assert_called_once()
                mock_imencode.assert_not_called()
        finally:
            import shutil
            shutil.rmtree(tmp_dir, ignore_errors=True)


class TestNextRunRetryAcceptance(unittest.TestCase):
    """Verify that the factory run-loop correctly retries failed uploads without changing SKUs."""
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.factory_dir = Path(self.tmp) / "FACTORY"
        self.sablon_dir = Path(self.tmp) / "SABLON"
        self.out_dir = Path(self.tmp) / "OUT"
        self.db_path = Path(self.tmp) / "state.db"
        
        self.factory_dir.mkdir()
        self.sablon_dir.mkdir()
        self.out_dir.mkdir()
        
        # Create a fake master and manifest using a real category name like "01_Wifi_Kutusu"
        manifest = {
            "batch_id": "test_batch",
            "items": [
                {
                    "master_id": "M01",
                    "theme_id": "T01",
                    "theme_title": "Test Theme",
                    "style_key": "Style",
                    "horizontal_path": "M01_H.jpg",
                    "vertical_path": "M01_V.jpg"
                }
            ]
        }
        with open(self.factory_dir / "manifest.json", "w") as f:
            json.dump(manifest, f)
            
        # Create fake raw image
        import numpy as np
        import cv2
        cv2.imwrite(str(self.factory_dir / "M01_H.jpg"), np.zeros((10,10,3), dtype=np.uint8))
        cv2.imwrite(str(self.factory_dir / "M01_V.jpg"), np.zeros((10,10,3), dtype=np.uint8))
        
        # Create fake category template for 01_Wifi_Kutusu (which requires Yatay orientation)
        cat_dir = self.sablon_dir / "01_Wifi_Kutusu"
        cat_dir.mkdir()
        cv2.imwrite(str(cat_dir / "tmpl1.png"), np.zeros((10,10,3), dtype=np.uint8))

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_seo_title_is_stable_for_same_product_identity(self):
        from SISTEM_MOTORU.pc_master_factory import SeoEngine

        seo = SeoEngine()
        first = seo.generate_title(
            "01_Wifi_Kutusu",
            "Tema_A",
            identity_seed="sg-M00001",
        )

        # Disturb global randomness to prove title generation is independent.
        import random
        for _ in range(10):
            random.random()

        second = seo.generate_title(
            "01_Wifi_Kutusu",
            "Tema_A",
            identity_seed="sg-M00001",
        )
        self.assertEqual(first, second)

    def test_factory_next_run_retry(self):
        from SISTEM_MOTORU.pc_master_factory import PCMasterFactory, StateManager
        
        def approved_pool(prod_key):
            if prod_key != "01_Wifi_Kutusu":
                return []
            return [(
                "tmpl1.png",
                self.sablon_dir / "01_Wifi_Kutusu" / "tmpl1.png",
                {"profile_orientation": "H", "nominal_aspect": 4.0 / 3.0},
            )]

        with patch("SISTEM_MOTORU.pc_master_factory.FACTORY_DIR", new=self.factory_dir), \
             patch("SISTEM_MOTORU.pc_master_factory.SABLON_DIR", new=self.sablon_dir), \
             patch("SISTEM_MOTORU.pc_master_factory.OUTPUT_DIR", new=self.out_dir), \
             patch("SISTEM_MOTORU.pc_master_factory.DB_PATH", new=self.db_path), \
             patch("SISTEM_MOTORU.pc_master_factory.EXCEL_DIR", new=Path(self.tmp)), \
             patch("SISTEM_MOTORU.pc_master_factory.MockupRenderer.process_mockup") as mock_render, \
             patch("SISTEM_MOTORU.pc_master_factory.AsyncCloudUploader.upload_sync") as mock_upload, \
             patch.object(
                 PCMasterFactory,
                 "get_template_pool",
                 side_effect=approved_pool,
             ):
             
            def fake_submit(fn, *args, **kwargs):
                class FakeFuture:
                    def result(self): pass
                fn(*args, **kwargs)
                return FakeFuture()
                
            with patch("SISTEM_MOTORU.pc_master_factory.ThreadPoolExecutor.submit", side_effect=fake_submit):
                # Mocking to skip real CV2 ops but return success
                mock_render.return_value = True
                
                # First run: fail upload
                mock_upload.return_value = None
                
                # Run 1
                bot = PCMasterFactory(mode="dongusel")
                bot.process_all()
                
                # Check DB state
                sm = StateManager(db_path=self.db_path)
                
                print("\n--- DB DEBUG ---")
                for table in ["system_meta", "render_artifacts", "products"]:
                    print(f"Table {table}:", sm.conn.execute(f"SELECT * FROM {table}").fetchall())
                print("----------------\n")
                
                artifacts = sm.conn.execute("SELECT * FROM render_artifacts").fetchall()
                self.assertEqual(len(artifacts), 1)
                self.assertEqual(artifacts[0]["status"], "UPLOAD_FAILED")
                
                # Capture SKU from system_meta
                meta_run1 = sm.conn.execute("SELECT value FROM system_meta WHERE key LIKE 'sku_%'").fetchall()
                self.assertEqual(len(meta_run1), 1)
                sku1 = meta_run1[0]["value"]
                
                # Run 2: succeed upload
                mock_upload.return_value = "http://success.com/mock.jpg"
                bot2 = PCMasterFactory(mode="dongusel")
                bot2.process_all()
                
                # Check DB state after run 2
                sm2 = StateManager(db_path=self.db_path)
                artifacts_run2 = sm2.conn.execute("SELECT * FROM render_artifacts").fetchall()
                self.assertEqual(len(artifacts_run2), 1, "Should not create duplicate artifact")
                self.assertEqual(
                    artifacts_run2[0]["status"],
                    "COMPLETED",
                    "Successful retry must finish upload + Excel transaction",
                )
                
                meta_run2 = sm2.conn.execute("SELECT value FROM system_meta WHERE key LIKE 'sku_%'").fetchall()
                self.assertEqual(len(meta_run2), 1)
                sku2 = meta_run2[0]["value"]
                
                self.assertEqual(sku1, sku2, "SKU must not change on retry")

if __name__ == "__main__":
    unittest.main()
