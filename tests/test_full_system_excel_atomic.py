# -*- coding: utf-8 -*-
"""True factory Excel appender: no duplicates, no truncated workbook."""
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT=Path(__file__).resolve().parents[1]
CHILD=r'''
import hashlib,importlib,os,sys
from pathlib import Path
from unittest.mock import patch
output=Path(sys.argv[1])
sys.argv.append("--preflight")
sys.path.insert(0,str(Path.cwd()/"99_SISTEM_MOTORLARI_VE_KODLAR"))
m=importlib.import_module("SISTEM_MOTORU.pc_master_factory")
a=m.FastExcelManager(output)
b=m.FastExcelManager(output)
row={"Stok Kodu":"S1","Mockup Dosya Adı":"MOCK_001.jpg","Ürün Adı":"Sentetik"}
a.append_row(row)
b.append_row(row)
assert a.flush() and b.flush()
import openpyxl
wb=openpyxl.load_workbook(output,read_only=True)
assert wb.active.max_row==2,wb.active.max_row
wb.close()
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
before=sha(output)
a.append_row({"Stok Kodu":"S2","Mockup Dosya Adı":"MOCK_002.jpg"})
with patch.object(m.os,"replace",side_effect=OSError("SIMULATED_EXCEL_REPLACE_FAILURE")):
    assert a.flush() is False
assert sha(output)==before
assert list((output.parent/"_EXCEL_ONCEKI_SURUMLER").glob("*.xlsx"))
assert a.flush() is True
wb=openpyxl.load_workbook(output,read_only=True)
assert wb.active.max_row==3
wb.close()
print("EXCEL_VERSIONED_ATOMIC_AND_DEDUP_PASS")
'''

def test_excel_flush_atomic_snapshot_and_multiprocess_stale_index():
    with tempfile.TemporaryDirectory(prefix="p0-excel-fake-") as td:
        output=Path(td)/"uploads.xlsx"
        env=os.environ.copy()
        env.update({"MOCKUP_RUNTIME_WORKSPACE":str(Path(td)/"fake-runtime"),"PYTHONUTF8":"1",
                    "PYTHONDONTWRITEBYTECODE":"1"})
        p=subprocess.run([sys.executable,"-B","-c",CHILD,str(output)],cwd=str(ROOT),
                         env=env,capture_output=True,text=True,encoding="utf-8",
                         errors="replace",timeout=90)
        assert p.returncode==0,p.stdout+"\n"+p.stderr
        assert "EXCEL_VERSIONED_ATOMIC_AND_DEDUP_PASS" in p.stdout
