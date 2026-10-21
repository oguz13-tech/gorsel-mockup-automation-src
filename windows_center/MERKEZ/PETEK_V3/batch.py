"""PETEK V3 durable local-only serial art rendering and Excel/CSV export."""
from __future__ import annotations
import contextlib
import csv
import hashlib
import json
import os
import logging
from pathlib import Path
import re
import sqlite3
import threading
from openpyxl import Workbook,load_workbook
from . import engine
from .registry import Registry

EXT={'.jpg','.jpeg','.png','.webp'}
HEADERS=['Stok Kodu','Dahili Barkod (GS1 Degil)','Model Kodu','Kategori','Urun Adi',
         'Tasarim Dosyasi','Sablon','Mockup Dosyasi','Yerel Dosya','Islem Durumu','Gorsel SHA256']


def _escape(text):
    data=str(text or '')
    return "'"+data if data.lstrip().startswith(('=','+','-','@')) else data


def iter_designs(folder):
    for directory,dirs,files in os.walk(folder,followlinks=False):
        dirs[:]=sorted([name for name in dirs if not (Path(directory)/name).is_symlink()])
        for name in sorted(files):
            p=Path(directory)/name
            if p.suffix.lower() in EXT and not p.is_symlink():yield p


class Ledger:
    def __init__(self,root):
        self.file=Path(root)/'petek_uretim.sqlite3'
        if self.file.is_symlink():raise engine.MaskError('Kayit DB baglantisi guvensiz.')
        with self.connect() as con:
            con.execute('''CREATE TABLE IF NOT EXISTS artifact (
              key TEXT PRIMARY KEY, sku TEXT NOT NULL, art_sha TEXT NOT NULL,
              template_sha TEXT NOT NULL,mask_sha TEXT NOT NULL,
              art_name TEXT NOT NULL, uid TEXT NOT NULL, output TEXT NOT NULL,
              output_sha TEXT NOT NULL DEFAULT '', status TEXT NOT NULL,
              message TEXT NOT NULL DEFAULT '', updated TEXT DEFAULT CURRENT_TIMESTAMP)''')
            con.execute('CREATE INDEX IF NOT EXISTS artifact_status_idx ON artifact(status)')

    def connect(self):
        con=sqlite3.connect(self.file,timeout=30)
        con.row_factory=sqlite3.Row
        con.execute('PRAGMA busy_timeout=30000')
        return con

    def reserve(self,fields):
        with self.connect() as con:
            con.execute('''INSERT OR IGNORE INTO artifact
              (key,sku,art_sha,template_sha,mask_sha,art_name,uid,output,status)
              VALUES (:key,:sku,:art_sha,:template_sha,:mask_sha,:art_name,:uid,:output,'PENDING')''',fields)
            found=dict(con.execute('SELECT * FROM artifact WHERE key=?',(fields['key'],)).fetchone())
        if any(found[k]!=fields[k] for k in fields):raise engine.MaskError('Kayit kimligi eski islemle cakisiyor.')
        return found

    def complete(self,key,sha):
        with self.connect() as con:
            c=con.execute('''UPDATE artifact SET output_sha=?,status='RENDERED',updated=CURRENT_TIMESTAMP
                   WHERE key=? AND status='PENDING' ''',(sha,key))
            if c.rowcount!=1:raise engine.MaskError('URETIM_KAYDI_TUTARSIZ')

    def quarantine(self,key,error):
        with self.connect() as con:
            con.execute('''UPDATE artifact SET status='REVIEW_REQUIRED',message=?,updated=CURRENT_TIMESTAMP
                        WHERE key=? AND status!='RENDERED' ''',(str(error)[:300],key))

    def successful(self):
        with self.connect() as con:
            return [dict(row) for row in con.execute('SELECT * FROM artifact WHERE status=\'RENDERED\' ORDER BY sku,uid')]


class Batch:
    def __init__(self,registry=None):
        self.registry=registry or Registry()
        self.ledger=Ledger(self.registry.root)
        self.cancel=threading.Event()

    def run(self,design_folder,template_uids,group='LOCAL',notify=None,max_jobs=5000,placement='cover',lighting='natural'):
        log=logging.getLogger('petek_v3')
        notify=notify or (lambda message:None)
        log.info('URETIM_BASLADI grubu=%s template_adedi=%s yerlestirme=%s',group,len(template_uids),placement)
        if placement not in engine.PLACEMENT_MODES:
            raise engine.MaskError('Yerlesim modu gecersiz: cover veya contain.')
        if not re.fullmatch('[A-Za-z0-9_-]{2,36}',str(group)):
            raise engine.MaskError('Magaza grubu yalnizca harf, sayi, alt cizgi, tire icerebilir.')
        if not 1<=int(max_jobs)<=10000:raise engine.MaskError('Is limiti 1-10000 arasinda olmali.')
        designs=Path(design_folder).resolve(strict=True)
        if not designs.is_dir():raise engine.MaskError('Tasarim klasoru yok.')
        private=self.registry.root.resolve()
        if designs==private or private in designs.parents or designs in private.parents:
            raise engine.MaskError('Tasarim klasoru cikti/kayit dizini icinde olamaz.')
        approved=[]
        for uid in list(dict.fromkeys(template_uids)):
            p,m=self.registry.get(uid,require_approved=True)
            source=Path(p['source']).resolve(strict=True)
            if source==designs or designs in source.parents or source in designs.parents:
                raise engine.MaskError('Tasarimlar kaynak PETEK sablonlarinin klasorunde olamaz.')
            approved.append((p,m))
        if not approved:raise engine.MaskError('Onayli PETEK sablonu secilmedi.')
        lock=self.registry.root/'PETEK_V3_FACTORY.lock'
        if lock.exists() or lock.is_symlink():raise engine.MaskError('Baska PETEK uretimi calisiyor / eski kilit insan kontrolu gerektiriyor.')
        try:
            fd=os.open(str(lock),os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
        except FileExistsError as exc:
            raise engine.MaskError('PETEK uretim kilidi dolu.') from exc
        with os.fdopen(fd,'w') as f:f.write(str(os.getpid()))
        stats={'status':'RUNNING','placement':placement,'attempted':0,'rendered':0,'reused':0,'rejected':0,'errors':[],'local_only':True,
               'marketplace_published':False}
        try:
            for art in iter_designs(designs):
                if self.cancel.is_set() or stats['attempted']>=max_jobs:break
                art=engine.safe_file(art)
                art_sha=engine.digest(art)
                for p,mask in approved:
                    if self.cancel.is_set() or stats['attempted']>=max_jobs:break
                    stats['attempted']+=1
                    name=p['uid']
                    source_sha=p['source_sha']
                    geometry_stamp=hashlib.sha256(json.dumps(
                        {'quad':p['quad'],'physical_aspect':p['physical_aspect']},
                        sort_keys=True,separators=(',', ':')).encode('utf-8')).hexdigest()
                    # V3.7 product code includes template, mask and source identity.
                    # Older ledgers/old SKU rows remain untouched. This prevents
                    # one artwork applied to two mockups sharing a product SKU.
                    product_hash=hashlib.sha256('|'.join((engine.RECIPE,placement,lighting,art_sha,
                         source_sha,p['mask_sha'],geometry_stamp,name,group)).encode('utf-8')).hexdigest()[:18].upper()
                    code=f'PDK-{group}-{product_hash}'
                    key=hashlib.sha256('|'.join((engine.RECIPE,placement,lighting,art_sha,source_sha,p['mask_sha'],
                                                geometry_stamp,name,group)).encode()).hexdigest()[:24]
                    output=self.registry.output/group/(code+'_'+key+'.png')
                    fields=dict(key=key,sku=code,art_sha=art_sha,template_sha=source_sha,
                                mask_sha=p['mask_sha'],art_name=art.name,uid=name,output=str(output))
                    try:
                        record=self.ledger.reserve(fields)
                        if record['status']=='RENDERED':
                            if not output.is_file() or engine.digest(output)!=record['output_sha']:
                                raise engine.MaskError('Onceden tamamlanan cikti bulunamadi veya degistirilmis.')
                            stats['reused']+=1;notify(f'ATLANDI / mevcut kayit: {name} / {art.name}');continue
                        if record['status']!='PENDING' or output.exists():
                            raise engine.MaskError('Supheli eski/artik tamamlanmamis cikti: inceleme gerek.')
                        # Source identity is revalidated after artwork read and
                        # immediately prior to rendering.
                        if engine.digest(p['source'])!=p['source_sha']:
                            raise engine.MaskError('Orijinal sablon degistirildi.')
                        picture,qa=engine.render(p['source'],art,p['quad'],mask,p['physical_aspect'],
                                                 marketplace=True,placement=placement,lighting=lighting)
                        image=engine.encode_png(picture)
                        engine.atomic_new_file(output,image)
                        receipt=engine.digest(output)
                        if engine.digest(p['source'])!=p['source_sha'] or engine.digest(art)!=art_sha:
                            raise engine.MaskError('Kaynak dosya uretim sirasinda degismis; cikti incelemeye alinmali.')
                        self.ledger.complete(key,receipt)
                        stats['rendered']+=1
                        log.info('URETILDI uid=%s art_sha=%s output_sha=%s',name,art_sha,receipt)
                        notify(f'URETILDI {stats["rendered"]}: {name} / {art.name} — delik/kavis orijinal korundu')
                    except (engine.MaskError,OSError,ValueError) as exc:
                        self.ledger.quarantine(key,str(exc))
                        stats['rejected']+=1
                        stats['errors'].append({'template':name,'design':art.name,'reason':str(exc)})
                        log.warning('ATLANDI uid=%s art_sha=%s neden=%s',name,art_sha,str(exc))
                        notify(f'ATLANDI: {name}/{art.name} — {exc}')
            csv_path,xlsx_path,count=self.export()
            stats.update(status='STOPPED' if self.cancel.is_set() else ('PARTIAL' if stats['rejected'] else 'LOCAL_READY'),
                         csv=str(csv_path),xlsx=str(xlsx_path),total_registry_outputs=count)
            log.info('URETIM_SONUCU %s',json.dumps({k:v for k,v in stats.items() if k!='errors'},ensure_ascii=False))
            return stats
        finally:
            lock.unlink(missing_ok=True)

    def export(self):
        rows=self.ledger.successful()
        fingerprint=hashlib.sha256('|'.join(row['key']+row['output_sha'] for row in rows).encode()).hexdigest()[:18]
        csv_path=self.registry.reports/f'PETEK_LISTE_{fingerprint}.csv'
        xlsx_path=self.registry.reports/f'PETEK_LISTE_{fingerprint}.xlsx'
        def entry(row):
            sku=row['sku']
            return [_escape(x) for x in (sku,'DAHILI-'+sku,'M-'+sku,'06_Petek_Dolabi',
                'Dekoratif Ahsap Petek Dolabi - '+Path(row['art_name']).stem,row['art_name'],row['uid'],
                Path(row['output']).name,row['output'],'YEREL_HAZIR',row['output_sha'])]
        for row in rows:
            path=Path(row['output'])
            if path.is_symlink() or not path.is_file() or engine.digest(path)!=row['output_sha']:
                raise engine.MaskError('Ticari liste icin cikti SHA dogrulanamadi: '+row['key'])
        if not csv_path.exists():
            from io import StringIO
            text=StringIO(newline='')
            writer=csv.writer(text);writer.writerow(HEADERS)
            for row in rows:writer.writerow(entry(row))
            engine.atomic_new_file(csv_path,('\ufeff'+text.getvalue()).encode('utf-8'))
        if not xlsx_path.exists():
            wb=Workbook(write_only=True)
            ws=wb.create_sheet('PETEK_YEREL');ws.append(HEADERS)
            for row in rows:ws.append(entry(row))
            import tempfile
            fd,tmp=tempfile.mkstemp(prefix='petek-',suffix='.xlsx',dir=str(self.registry.reports));os.close(fd)
            try:
                wb.save(tmp)
                engine.atomic_new_file(xlsx_path,Path(tmp).read_bytes())
            finally:Path(tmp).unlink(missing_ok=True)
        # Existing files must be valid snapshots, not silent name collisions.
        with csv_path.open('r',encoding='utf-8-sig',newline='') as f:
            reader=csv.reader(f)
            if next(reader,None)!=HEADERS or list(reader)!=[entry(r) for r in rows]:
                raise engine.MaskError('CSV export beklenen icerikle uyusmuyor.')
        wb=load_workbook(xlsx_path,read_only=True,data_only=True)
        try:
            if [list(row) for row in wb.active.iter_rows(values_only=True)] != [HEADERS]+[entry(r) for r in rows]:
                raise engine.MaskError('EXCEL export beklenen icerikle uyusmuyor.')
        finally:wb.close()
        return csv_path,xlsx_path,len(rows)
