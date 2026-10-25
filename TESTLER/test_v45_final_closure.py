"""New V4.5 missing-feature regressions: cooperative pause, RO operations,
commerce-template export, and large-image lossless source proxy safeguards."""
import os
import sys
from pathlib import Path
import sqlite3
import threading
import time

import cv2
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "MERKEZ"))

from URETIM_V4.unified_batch import UnifiedV4Batch
from URETIM_V4.ops_dashboard_v45 import summarize
from URETIM_V4.marketplace_excel_v45 import detect_columns, export_from_template
from URETIM_V4.design_proxy_v43 import _large_png_or_webp_proxy
from PETEK_V3 import engine as petek_engine


def _ledger(root: Path, path: Path, sha: str, status: str = 'RENDERED'):
    root.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(root / 'v4_ortak_uretim.sqlite3')
    conn.execute('CREATE TABLE artifact(key TEXT PRIMARY KEY,sku TEXT,category TEXT,art_sha TEXT,template_sha TEXT,mask_sha TEXT,art_name TEXT,uid TEXT,output TEXT,output_sha TEXT,status TEXT)')
    conn.execute('INSERT INTO artifact VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                 ('k1','SKU_1','03_Ahsap_Tablo','art','source','mask','design.png','uid',str(path),sha,status))
    conn.commit();conn.close()
    return root / 'v4_ortak_uretim.sqlite3'


def _output(tmp_path):
    root = tmp_path / 'state'
    folder = root / 'uretilen_gorseller' / 'LOCAL' / '03_Ahsap_Tablo'
    folder.mkdir(parents=True)
    path = folder / 'mockup.png'
    path.write_bytes(b'known output bytes')
    return root, path, petek_engine.digest(path)


def test_pause_resume_finish_between_jobs():
    obj=object.__new__(UnifiedV4Batch)
    obj.pause=threading.Event();obj.cancel=threading.Event()
    obj.request_pause()
    events=[]
    out=[]
    entered=threading.Event()
    def announce(text):
        events.append(text)
        entered.set()
    worker=threading.Thread(target=lambda:out.append(obj._wait_if_paused(announce)))
    worker.start();assert entered.wait(2.0)
    assert worker.is_alive() and not out
    obj.resume();worker.join(timeout=2)
    assert out == [True] and not obj.pause.is_set()
    assert events and 'DURAKLATILDI' in events[0]


def test_stop_works_while_paused():
    obj=object.__new__(UnifiedV4Batch)
    obj.pause=threading.Event();obj.cancel=threading.Event()
    obj.request_pause();obj.request_stop()
    assert obj.cancel.is_set() and not obj.pause.is_set()
    assert obj._wait_if_paused(lambda _:None) is False


def test_read_only_summary_matches_file_sha(tmp_path):
    root,p,sha=_output(tmp_path)
    db=_ledger(root,p,sha)
    before=petek_engine.digest(db)
    report=summarize(root)
    assert report['verified']==1 and report['total']==1
    assert report['categories']['03_Ahsap_Tablo']==1
    assert petek_engine.digest(db)==before


def test_missing_and_mismatch_summary(tmp_path):
    root,p,sha=_output(tmp_path)
    _ledger(root,p,sha)
    p.write_bytes(b'wrong')
    assert summarize(root)['sha_mismatch']==1
    p.unlink()
    assert summarize(root)['missing']==1


def test_dashboard_never_creates_missing_db(tmp_path):
    state=tmp_path/'state';state.mkdir()
    with pytest.raises(ValueError):summarize(state)
    assert list(state.iterdir())==[]


def test_marketplace_aliases():
    m=detect_columns(['Satıcı Stok Kodu', 'Görsel 1', 'Barkod', 'Ürün Açıklaması'])
    assert m['sku']==1 and m['image_url']==2 and m['barcode']==3 and m['description']==4


def test_marketplace_rejects_unknown_headers():
    with pytest.raises(ValueError):detect_columns(['ABC','DEF'])
    assert detect_columns(['sku', 'link do fotograf'],custom={'image_url':'link do fotograf'})['image_url']==2


def _cloud_db(root, sha):
    db=sqlite3.connect(root/'cloud_urls_v43.sqlite3')
    db.execute('CREATE TABLE cloud_link(artifact_key TEXT PRIMARY KEY,output_sha TEXT,url TEXT,recorded_at TEXT)')
    db.execute('INSERT INTO cloud_link VALUES(?,?,?,?)',('k1',sha,'https://images.example.com/mockup.png','now'))
    db.commit();db.close()


def test_excel_actual_https_export_new_file_only(tmp_path):
    root,p,sha=_output(tmp_path);_ledger(root,p,sha);_cloud_db(root,sha)
    wb=Workbook();ws=wb.active
    ws.append(['Satıcı Stok Kodu','Görsel 1','Barkod','Ürün Adı'])
    template=tmp_path/'empty.xlsx';wb.save(template);wb.close()
    before=petek_engine.digest(template)
    result=export_from_template(template,root,[{'key':'k1'}],report_root=root/'raporlar')
    assert result['rows']==1 and result['status']=='VERIFY_PRODUCT_FIELDS_MANUALLY'
    assert petek_engine.digest(template)==before
    wb_out=load_workbook(result['output'])
    try:
        ws=wb_out.active
        assert ws['A2'].value=='SKU_1'
        assert ws['B2'].value=='https://images.example.com/mockup.png'
        assert ws['B2'].hyperlink.target=='https://images.example.com/mockup.png'
        assert ws['C2'].value is None and ws['D2'].value is None
    finally:
        wb_out.close()
    # idempotent exact same output path; no overwrite
    result2=export_from_template(template,root,[{'key':'k1'}],report_root=root/'raporlar')
    assert result['output']==result2['output']


def test_excel_input_with_existing_product_is_rejected(tmp_path):
    root,p,sha=_output(tmp_path);_ledger(root,p,sha);_cloud_db(root,sha)
    wb=Workbook();ws=wb.active;ws.append(['Stok Kodu','Resim Linki']);ws.append(['EXISTING','https://old.example/img'])
    template=tmp_path/'has_data.xlsx';wb.save(template);wb.close()
    with pytest.raises(ValueError,match='mevcut ürün'):
        export_from_template(template,root,[{'key':'k1'}],report_root=root/'raporlar')


def test_excel_link_sha_mismatch_not_exported(tmp_path):
    root,p,sha=_output(tmp_path);_ledger(root,p,sha);_cloud_db(root,'f'*64)
    wb=Workbook();wb.active.append(['Stok Kodu','Resim Linki'])
    t=tmp_path/'empty.xlsx';wb.save(t);wb.close()
    with pytest.raises(ValueError,match='SHA doğrulanmış HTTPS'):
        export_from_template(t,root,[{'key':'k1'}],report_root=root/'raporlar')


def test_large_proxy_requires_explicit_magick(tmp_path,monkeypatch):
    src=tmp_path/'huge.png';src.write_bytes(b'xyz')
    monkeypatch.delenv('MAGICK_EXE',raising=False)
    with pytest.raises(petek_engine.MaskError,match='MAGICK_EXE'):
        _large_png_or_webp_proxy(src,petek_engine.digest(src),tmp_path/'working')
    assert src.read_bytes()==b'xyz'


def test_magick_source_preserved_and_proxy_created(tmp_path,monkeypatch):
    executable=Path('/opt/imagemagick/bin/magick')
    if not executable.is_file():pytest.skip('Only on ImageMagick test runner')
    monkeypatch.setenv('MAGICK_EXE',str(executable))
    src=tmp_path/'source.png';pixels=np.ones((256,480,3),np.uint8)*128
    assert cv2.imwrite(str(src),pixels)
    sha=petek_engine.digest(src)
    proxy=_large_png_or_webp_proxy(src,sha,tmp_path/'working')
    assert proxy.is_file() and petek_engine.digest(src)==sha
    assert cv2.imread(str(proxy)).shape==pixels.shape
    assert _large_png_or_webp_proxy(src,sha,tmp_path/'working')==proxy


def test_marketplace_header_after_instructions(tmp_path):
    root,p,sha=_output(tmp_path);_ledger(root,p,sha);_cloud_db(root,sha)
    wb=Workbook();ws=wb.active
    ws.append(['E-ticaret Excel Aktarım Şablonu'])
    ws.append(['Ürünler aşağıdaki başlıklara yazılacak'])
    ws.append(['Stok Kodu','Görsel 1'])
    original=tmp_path/'instructions.xlsx';wb.save(original);wb.close()
    result=export_from_template(original,root,[{'key':'k1'}],report_root=root/'raporlar')
    assert result['header_row']==3
    exported=load_workbook(result['output'])
    try:
        assert exported.active['A4'].value=='SKU_1'
        assert exported.active['B4'].hyperlink.target=='https://images.example.com/mockup.png'
        assert exported.active['A1'].value=='E-ticaret Excel Aktarım Şablonu'
    finally:
        exported.close()
