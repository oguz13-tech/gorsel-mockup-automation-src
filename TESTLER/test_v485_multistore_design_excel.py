"""V4.8.5 user-locked store identities and NO-extra-Excel contract tests.

Entirely synthetic and offline, no network calls and no user file access.
"""
from pathlib import Path
import sqlite3
import sys
import pytest
from openpyxl import Workbook,load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'MERKEZ'))
from URETIM_V4.store_identity_v485 import (store_codes,codes_from_sku,
                                            filename,design_number)
from URETIM_V4.store_profiles_v47 import StoreProfiles
from URETIM_V4.unified_batch import UnifiedV4Ledger
from URETIM_V4.cloud_excel_v43 import CloudExcel
from URETIM_V4.store_excel_v485 import export_store_upload
from PETEK_V3 import engine


def test_same_print_design_number_in_five_stores_with_different_mockups():
    sources=[]
    for index in range(1,6):
        store=f'M{index}'
        code=store_codes('18472',store,'01_Wifi_Kutusu')
        assert code['sku']==f'18472-WifiKutusu-{store}'
        assert code['barcode']==code['sku']
        assert code['model']==f'18472-{store}-WifiKutusu'
        assert codes_from_sku(code['sku'])==code
        mockup=filename('18472',store,'01_Wifi_Kutusu',f'Wifi_Kutusu_{35+index}')
        assert mockup==f'18472__{store}__WifiKutusu__TPL{35+index:03d}__V01.png'
        sources.append(mockup)
    assert len(set(sources))==5
    assert design_number('18472.png')=='18472'
    assert design_number('00018472.png')=='00018472'
    assert filename('18472','M5','01_Wifi_Kutusu','Wifi_35',extension='.jpg')==(
            '18472__M5__WifiKutusu__TPL035__V01.jpg')
    with pytest.raises(ValueError):design_number('18472 tasarım.png')
    with pytest.raises(ValueError):store_codes('18472','M5','UNSUPPORTED')


def _receipt(ledger,cloud,state,store,art,uid,number):
    category='01_Wifi_Kutusu'
    ident=store_codes(number,store,category)
    folder=state/'uretilen_gorseller'/store/category
    folder.mkdir(parents=True,exist_ok=True)
    path=folder/filename(number,store,category,uid)
    path.write_bytes(b'synthetic fixture '+store.encode()+uid.encode())
    fields={
        'key':store+'-'+uid+'-'+number,'sku':ident['sku'],'category':category,
        'art_sha':'a'*64,'template_sha':'b'*64,'mask_sha':'c'*64,
        'art_name':art,'uid':uid,'output':str(path),
    }
    ledger.reserve(fields)
    digest=engine.digest(path)
    ledger.complete(fields['key'],digest)
    with cloud.connect() as conn:
        conn.execute('INSERT INTO cloud_link(artifact_key,output_sha,url) VALUES(?,?,?)',
            (fields['key'],digest,'https://images.example.org/'+path.name))
    return fields


def _blank_template(path,multiple=True):
    wb=Workbook();sh=wb.active;sh.title='Ürün Yükleme'
    labels=['Stok Kodu','Barkod','Model Kodu','Kategori','Görsel 1']
    if multiple:labels.extend(['Görsel 2','Görsel 3'])
    sh.append(labels)
    wb.save(path)
    wb.close()


def test_selected_store_excel_not_extra_report_and_one_row_many_mockups(tmp_path):
    state=tmp_path/'GorselMockup';state.mkdir()
    ledger=UnifiedV4Ledger(state)
    cloud=CloudExcel(state,state/'raporlar')
    _receipt(ledger,cloud,state,'M5','18472.png','Wifi_Kutusu_035','18472')
    _receipt(ledger,cloud,state,'M5','18472.png','Wifi_Kutusu_036','18472')
    _receipt(ledger,cloud,state,'M4','18472.png','Wifi_Kutusu_044','18472')
    original=tmp_path/'blank.xlsx'
    _blank_template(original)
    original_bytes=original.read_bytes()
    # Cloud URLs for new store codes already exist: no second generic xlsx/csv.
    csv_path,xlsx_path,count=cloud.export(ledger.successful())
    assert csv_path is None and xlsx_path is None and count==3
    assert list((state/'raporlar').glob('*.xlsx'))==[]
    out5=export_store_upload(original,state,store_id='M5',report_root=state/'raporlar')
    out4=export_store_upload(original,state,store_id='M4',report_root=state/'raporlar')
    assert out5['rows']==1 and out5['images']==2
    assert out4['rows']==1 and out4['images']==1
    assert out5['output']!=out4['output']
    assert original.read_bytes()==original_bytes
    wb=load_workbook(out5['output'])
    sheet=wb.active
    assert sheet.cell(2,1).value=='18472-WifiKutusu-M5'
    assert sheet.cell(2,2).value=='18472-WifiKutusu-M5'
    assert sheet.cell(2,3).value=='18472-M5-WifiKutusu'
    assert sheet.cell(2,5).hyperlink.target.startswith('https://')
    assert sheet.cell(2,6).hyperlink.target.startswith('https://')
    assert sheet.max_row==2
    wb.close()
    wb=load_workbook(out4['output'])
    assert wb.active.cell(2,1).value=='18472-WifiKutusu-M4'
    assert wb.active.cell(2,6).value is None
    wb.close()


def test_store_excel_never_silently_drops_extra_images(tmp_path):
    state=tmp_path/'state';state.mkdir()
    ledger=UnifiedV4Ledger(state)
    cloud=CloudExcel(state,state/'raporlar')
    _receipt(ledger,cloud,state,'M5','18472.png','Wifi_035','18472')
    _receipt(ledger,cloud,state,'M5','18472.png','Wifi_036','18472')
    original=tmp_path/'one_image.xlsx';_blank_template(original,multiple=False)
    with pytest.raises(ValueError,match='SUTUNLARI_YETERSIZ'):
        export_store_upload(original,state,store_id='M5',report_root=state/'raporlar')
    assert list((state/'raporlar').glob('*.xlsx'))==[]


def test_store_excel_preserves_nonempty_template_and_refuses_wrong_url(tmp_path):
    state=tmp_path/'state';state.mkdir()
    ledger=UnifiedV4Ledger(state)
    cloud=CloudExcel(state,state/'raporlar')
    _receipt(ledger,cloud,state,'M5','18472.png','Wifi_035','18472')
    template=tmp_path/'template.xlsx';_blank_template(template)
    wb=load_workbook(template);wb.active.cell(2,1,'EXISTING PRODUCT');wb.save(template);wb.close()
    with pytest.raises(ValueError,match='DOLU_SABLON'):
        export_store_upload(template,state,store_id='M5',report_root=state/'raporlar')
    assert template.exists()


def test_store_assignment_no_phash_image_decoding(tmp_path,monkeypatch):
    manager=StoreProfiles(tmp_path/'state')
    manager.add_store('M4');manager.add_store('M5')
    a={'category':'01_Wifi_Kutusu','uid':'Wifi_035','source_sha':'a'*64}
    b={'category':'01_Wifi_Kutusu','uid':'Wifi_036','source_sha':'b'*64}
    import URETIM_V4.store_profiles_v47 as module
    monkeypatch.setattr(module,'perceptual_fingerprint',lambda _: (_ for _ in ()).throw(
        AssertionError('NO pHash checks requested by operator')))
    manager.assign('M4',[a],all_approved=[a,b])
    manager.assign('M5',[b],all_approved=[a,b])
    with pytest.raises(ValueError,match='başka|REUSED_ACROSS_STORES'):
        manager.assign('M4',[b],all_approved=[a,b])
