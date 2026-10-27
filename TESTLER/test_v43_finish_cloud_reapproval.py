from pathlib import Path
import hashlib
import sys
import sqlite3

import cv2
import numpy as np
from openpyxl import load_workbook
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'MERKEZ'))
from URETIM_V4.finish_v42 import subdued_reflection
from URETIM_V4.cloud_excel_v43 import CloudExcel, REMOTE_HEADERS
from URETIM_V4.unified_batch import unified_receipt_identity, UnifiedV4Batch


def test_cam_reflection_visible_but_only_inside_mask():
    h,w=300,340
    base=np.full((h,w,3), (42,67,110),dtype=np.uint8)
    source=np.full((h,w,3), 125, dtype=np.uint8)
    quad=[[50,40],[291,40],[291,253],[50,253]]
    mask=np.zeros((h,w), np.uint8)
    cv2.fillConvexPoly(mask,np.array(quad,dtype=np.int32),255)
    on=subdued_reflection(source,base,mask,quad,'02_Cam_Tablo_Dikey','strong')
    off=subdued_reflection(source,base,mask,quad,'02_Cam_Tablo_Dikey','off')
    delta=np.abs(on.astype(np.int16)-off.astype(np.int16)).mean(axis=2)
    assert np.array_equal(on[mask==0],base[mask==0])
    assert np.array_equal(off,base)
    assert float(np.percentile(delta[mask>0],70)) >= 4.0
    assert delta.max() <= 45
    notglass=subdued_reflection(source,base,mask,quad,'01_Wifi_Kutusu','strong')
    assert np.array_equal(notglass,base)


def test_editing_glass_effect_or_corners_yields_new_sku_identity():
    info=dict(category='02_Cam_Tablo_Yatay',uid='example',source_sha='a'*64,
              mask_sha='b'*64,quad=[[30,30],[200,30],[200,170],[30,170]],physical_aspect=1.5)
    a,_=unified_receipt_identity(info,'c'*64,group='LOCAL',placement='cover',lighting='natural',glass_effect='on',corner_radius=.008)
    b,_=unified_receipt_identity(info,'c'*64,group='LOCAL',placement='cover',lighting='natural',glass_effect='off',corner_radius=.008)
    c,_=unified_receipt_identity(info,'c'*64,group='LOCAL',placement='cover',lighting='natural',glass_effect='on',corner_radius=.015)
    d,_=unified_receipt_identity({**info,'quad':[[31,30],[200,30],[200,170],[30,170]]},'c'*64,group='LOCAL',placement='cover',lighting='natural',glass_effect='on',corner_radius=.008)
    assert len({a,b,c,d})==4


def make_row(tmp_path, state='localstate'):
    f=tmp_path/state/'uretilen_gorseller'/'LOCAL'/'04_Cam_Dograma'/'hello.png'
    f.parent.mkdir(parents=True,exist_ok=True)
    f.write_bytes(b'safe mockup rendered bytes')
    return dict(key='k001',sku='CMD-LOCAL-001',category='04_Cam_Dograma',
                art_name='art.png',uid='CMD_01',output=str(f),output_sha=hashlib.sha256(f.read_bytes()).hexdigest())


def test_old_outputs_are_published_once_then_excel_backfilled(tmp_path, monkeypatch):
    runtime=tmp_path/'runtime';runtime.mkdir()
    (runtime/'ayarlar.json').write_text('{"bulut_ayarlari":{"imgbb_keys":["fake-test-key"]}}')
    row=make_row(tmp_path)
    c=CloudExcel(tmp_path/'localstate', tmp_path/'localstate'/'raporlar',runtime)
    calls=[]
    def upload(p,keys,free):
        calls.append(p);return 'https://i.ibb.co/abc/preview.png'
    outcome=c.publish_existing([row],upload_fn=upload)
    assert outcome['linked']==1 and outcome['uploaded']==1 and len(calls)==1
    wb=load_workbook(outcome['xlsx'])
    assert wb.active['J2'].value=='https://i.ibb.co/abc/preview.png'
    assert wb.active['J2'].hyperlink.target=='https://i.ibb.co/abc/preview.png'
    assert wb.active['A2'].value==row['sku']
    assert wb.active['L2'].hyperlink.target.startswith('file:///')
    wb.close()
    first=Path(outcome['xlsx']).read_bytes()
    second=c.publish_existing([row],upload_fn=upload)
    assert second['reused']==1 and second['uploaded']==0 and len(calls)==1
    assert Path(outcome['xlsx']).read_bytes()==first


def test_cloud_bad_url_never_in_excel(tmp_path):
    runtime=tmp_path/'runtime';runtime.mkdir()
    (runtime/'ayarlar.json').write_text('{"bulut_ayarlari":{"imgbb_keys":["fake"]}}')
    row=make_row(tmp_path,state='s')
    c=CloudExcel(tmp_path/'s', tmp_path/'s'/'r', runtime)
    a=c.publish_existing([row],upload_fn=lambda *_: 'file:///C:/local.png')
    assert a['linked']==0 and a['failed']==1
    wb=load_workbook(a['xlsx']);assert wb.active.max_row==1;wb.close()


def test_local_safe_not_modify_original_approved_geometry_code():
    app=(ROOT/'SABLON_V2_MIMARI/app.py').read_text(encoding='utf-8')
    js=(ROOT/'SABLON_V2_MIMARI/static/js/main.js').read_text(encoding='utf-8')
    assert 'approved_geometry_revisions' in app
    assert 'SET surface_quad=?, orientation=?, nominal_aspect=?' in app
    assert 'glass_corner_radius' in js
    assert 'glass_corner_radius' in app


def test_unlimited_default_and_two_options_visible():
    import inspect
    from URETIM_V4.gui_v4 import LIGHT_LABELS
    assert len(LIGHT_LABELS)==2
    assert 'Kapalı' not in LIGHT_LABELS[0]
    assert 'max_jobs: int | None = None' in inspect.getsource(UnifiedV4Batch.run)


def test_huge_png_is_safely_rejected_before_full_decompression(tmp_path):
    import struct
    from URETIM_V4.design_proxy_v43 import prepare_render_artwork
    p=tmp_path/'huge.png'
    p.write_bytes(b'\x89PNG\r\n\x1a\n' + b'\x00\x00\x00\x0dIHDR' + struct.pack('>II', 24960,16640))
    with pytest.raises(Exception,match='415.3 MP'):
        prepare_render_artwork(p,tmp_path/'proxy')
    assert not (tmp_path/'proxy').exists()


def test_small_jpeg_not_rewritten(tmp_path):
    from URETIM_V4.design_proxy_v43 import prepare_render_artwork
    p=tmp_path/'small.jpg'
    cv2.imwrite(str(p),np.full((320,280,3),125,dtype=np.uint8))
    out,sha=prepare_render_artwork(p,tmp_path/'proxy')
    assert out==p
    assert sha==hashlib.sha256(p.read_bytes()).hexdigest()
    assert not (tmp_path/'proxy').exists()

def test_glass_reflection_also_visible_on_white_print():
    arr=np.full((310,310,3),255,np.uint8)
    mask=np.zeros((310,310),np.uint8)
    cv2.rectangle(mask,(40,30),(270,280),255,-1)
    q=[[40,30],[270,30],[270,280],[40,280]]
    on=subdued_reflection(arr,arr,mask,q,'05_Ocak_Arkasi','strong')
    assert np.array_equal(on[mask==0],arr[mask==0])
    assert np.percentile(np.abs(on.astype(float)-arr.astype(float))[mask>0],75)>=3