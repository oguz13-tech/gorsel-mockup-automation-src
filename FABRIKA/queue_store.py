"""Durable single-worker queue and local product/export ledger, no marketplace API."""
import contextlib
import csv
import io
import json
from pathlib import Path
import re
import sqlite3
import zipfile

from .safety import IntegrityError, immutable, json_new, key, read_json, safe, sha

CATEGORIES = {
    '01_Wifi_Kutusu':'WiFi Kutusu', '02_Cam_Tablo_Dikey':'Cam Tablo Dikey',
    '02_Cam_Tablo_Yatay':'Cam Tablo Yatay', '03_Ahsap_Tablo':'Ahşap Tablo',
    '04_Cam_Dograma':'Cam Doğrama', '05_Ocak_Arkasi':'Ocak Arkası',
    '06_Petek_Dolabi':'Petek Dolabı'}


class Store:
    def __init__(self, root):
        self.root = safe(root); self.root.mkdir(parents=True,exist_ok=True)
        self.db = safe(self.root/'factory_v3.sqlite3')
        with self.connect() as c:
            c.executescript('''
            CREATE TABLE IF NOT EXISTS products(
                number INTEGER PRIMARY KEY AUTOINCREMENT, product_key TEXT UNIQUE NOT NULL,
                category TEXT NOT NULL, art_sha TEXT NOT NULL, sku TEXT UNIQUE NOT NULL,
                barcode TEXT UNIQUE NOT NULL, model TEXT NOT NULL, title TEXT NOT NULL,
                description TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS jobs(
                job_key TEXT PRIMARY KEY, product_key TEXT NOT NULL REFERENCES products(product_key),
                payload TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'PENDING', attempts INTEGER NOT NULL DEFAULT 0,
                output TEXT NOT NULL DEFAULT '', output_sha TEXT NOT NULL DEFAULT '',
                quality TEXT NOT NULL DEFAULT '{}', error TEXT NOT NULL DEFAULT '');
            CREATE TABLE IF NOT EXISTS events(
                id INTEGER PRIMARY KEY AUTOINCREMENT, job_key TEXT NOT NULL,
                status TEXT NOT NULL, created TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP);
            ''')

    @contextlib.contextmanager
    def connect(self):
        c=sqlite3.connect(self.db, timeout=15); c.row_factory=sqlite3.Row
        c.execute('PRAGMA foreign_keys=ON'); c.execute('PRAGMA synchronous=FULL')
        try:
            with c: yield c
        finally: c.close()

    def product(self, category, art_sha, title):
        if category not in CATEGORIES: raise ValueError('Bilinmeyen kategori.')
        # Same PETEK v1 product key/PDK identity, independent of new render recipe.
        product_key=key({'category':category,'store':'LOCAL','art_sha':art_sha,
                         'recipe':'petek-product-v1' if category=='06_Petek_Dolabi' else 'factory-product-v1'})
        prefix='PDK' if category=='06_Petek_Dolabi' else 'GMF'
        sku=prefix+'-'+product_key[:16].upper(); barcode=prefix+product_key[:20].upper()
        title=re.sub(r'[_\-]+',' ',title).strip()[:65]+' '+CATEGORIES[category]
        description=CATEGORIES[category]+' için dekoratif tasarım. Ölçü ve malzeme bilgilerini satış öncesi doğrulayın.'
        with self.connect() as c:
            c.execute('INSERT OR IGNORE INTO products(product_key,category,art_sha,sku,barcode,model,title,description) VALUES(?,?,?,?,?,?,?,?)',
                (product_key,category,art_sha,sku,barcode,'M-'+sku,title,description))
            row=c.execute('SELECT * FROM products WHERE product_key=?',(product_key,)).fetchone()
            if row is None or row['sku']!=sku or row['barcode']!=barcode or row['art_sha']!=art_sha:
                raise IntegrityError('SKU/barkod çakışması; numara yeniden atanmadı.')
            return dict(row)

    def enqueue(self, identity, product, payload):
        job=key(identity)
        with self.connect() as c:
            old=c.execute('SELECT * FROM jobs WHERE job_key=?',(job,)).fetchone()
            if old and old['product_key']!=product['product_key']:raise IntegrityError('İş kimliği çakışması.')
            if old:
                # Fresh immutable input snapshot path, same logical content identity.
                c.execute('UPDATE jobs SET payload=? WHERE job_key=?',(json.dumps(payload,ensure_ascii=False),job))
            else:
                c.execute('INSERT INTO jobs(job_key,product_key,payload) VALUES(?,?,?)',(job,product['product_key'],json.dumps(payload,ensure_ascii=False)))
                c.execute('INSERT INTO events(job_key,status) VALUES(?,?)',(job,'PENDING'))
        return job

    def row(self, job):
        with self.connect() as c:return dict(c.execute('SELECT * FROM jobs WHERE job_key=?',(job,)).fetchone())

    def status(self, job, status, **values):
        allowed={'status','output','output_sha','quality','error','attempts'}
        if set(values)-allowed:raise ValueError('Geçersiz iş alanı.')
        values['status']=status
        with self.connect() as c:
            c.execute('UPDATE jobs SET '+','.join(k+'=?' for k in values)+' WHERE job_key=?',(*values.values(),job))
            c.execute('INSERT INTO events(job_key,status) VALUES(?,?)',(job,status))

    def reuse(self, job):
        row=self.row(job)
        if row['output']:
            p=safe(row['output'])
            if not p.is_file() or sha(p)!=row['output_sha']:
                raise IntegrityError('Önceki görsel kayıp/değişmiş; üzerine yazılmadı: '+str(p))
            return row
        return None

    def rendered(self, job, output, quality):
        self.status(job,'RENDERED',output=str(output),output_sha=sha(output),quality=json.dumps(quality,ensure_ascii=False),error='')
        from PIL import Image
        with Image.open(output) as im:
            im.verify()
        with Image.open(output) as im:
            if im.size!=(1200,1600):raise IntegrityError('Çıktı 1200×1600 değil.')
        self.status(job,'VERIFIED')

    def records(self, jobs):
        result=[]
        with self.connect() as c:
            for job in sorted(set(jobs)):
                row=c.execute('SELECT j.*,p.number,p.category,p.art_sha,p.sku,p.barcode,p.model,p.title,p.description FROM jobs j JOIN products p ON j.product_key=p.product_key WHERE job_key=?',(job,)).fetchone()
                if row and row['output']:result.append(dict(row))
        return result


FIELDS=['Ürün No','SKU','Dahili Barkod (GS1 değildir)','Model Kodu','Kategori','Ürün Adı','Açıklama','Görsel','Tasarım SHA256','İş Kimliği','Yerel Durum','Görsel Onayı']


def cell(value):
    text=str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@')) else text


def export(store, jobs):
    """Export only verified receipts; state moves after both immutable files verify."""
    from openpyxl import Workbook,load_workbook
    records=store.records(jobs)
    for r in records:
        store.reuse(r['job_key'])
        if r['status'] not in ('VERIFIED','EXPORTED','COMPLETED'):
            raise IntegrityError('Doğrulanmamış iş dışa aktarılamaz.')
    rows=[[r['number'],r['sku'],r['barcode'],r['model'],CATEGORIES[r['category']],r['title'],r['description'],r['output'],r['art_sha'],r['job_key'],'EXPORTED','İNSAN İNCELEMESİ BEKLİYOR'] for r in records]
    snapshot=key({'fields':FIELDS,'rows':rows}); folder=store.root/'EXPORTS'/snapshot
    folder.mkdir(parents=True,exist_ok=True)
    receipt=folder/'RECEIPT.json'; csv_path=folder/'URUNLER.csv'; xlsx_path=folder/'URUNLER.xlsx'
    if receipt.exists():
        meta=read_json(receipt)
        if meta.get('snapshot')!=snapshot or sha(csv_path)!=meta.get('csv_sha') or sha(xlsx_path)!=meta.get('xlsx_sha'):
            raise IntegrityError('Önceki Excel/CSV değişmiş; üzerine yazılmadı.')
    else:
        stream=io.StringIO(newline='');writer=csv.writer(stream);writer.writerow(FIELDS)
        for row in rows:writer.writerow([cell(v) for v in row])
        immutable(csv_path,stream.getvalue().encode('utf-8-sig'))
        wb=Workbook();ws=wb.active;ws.title='Yerel Ürünler';ws.append(FIELDS)
        for row in rows:ws.append([cell(v) for v in row])
        # Deterministic XLSX bytes allow recovery after export-before-receipt crash.
        from datetime import datetime
        wb.properties.created=wb.properties.modified=datetime(2000,1,1)
        buffer=io.BytesIO();wb.save(buffer);wb.close();canonical=io.BytesIO()
        with zipfile.ZipFile(buffer) as source,zipfile.ZipFile(canonical,'w',zipfile.ZIP_DEFLATED) as dest:
            for name in sorted(source.namelist()):
                data=source.read(name)
                if name=='docProps/core.xml':
                    data=re.sub(rb'<dcterms:modified[^>]*>.*?</dcterms:modified>',b'<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:modified>',data)
                info=zipfile.ZipInfo(name,(2000,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;dest.writestr(info,data)
        immutable(xlsx_path,canonical.getvalue())
        check=load_workbook(xlsx_path,read_only=True)
        try:
            if check.active.max_row!=len(rows)+1:raise IntegrityError('Excel satır doğrulaması başarısız.')
        finally:check.close()
        json_new(receipt,{'snapshot':snapshot,'csv_sha':sha(csv_path),'xlsx_sha':sha(xlsx_path),'rows':len(rows)})
    for r in records:
        if r['status']!='COMPLETED':
            store.status(r['job_key'],'EXPORTED');store.status(r['job_key'],'COMPLETED')
    return {'csv':str(csv_path),'excel':str(xlsx_path),'rows':len(rows),'human_review_required':True}
